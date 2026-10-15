"""The few things python-pptx has no public API for, when WRITING a deck.

Like reader.py on the import side, this is the ONLY export module that
touches raw XML (`_element`, `_p`, XPath, private parts).
"""
from __future__ import annotations

import copy
import uuid

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn

from ..reader import P14, paragraph_style

_SECTIONS_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
_SVG_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
_INDENT = 285750  # EMU per list level (0.3125 in)
_BU_TAGS = ("a:buClrTx", "a:buClr", "a:buSzTx", "a:buSzPct", "a:buSzPts", "a:buFontTx", "a:buFont",
            "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip")
_AFTER_BULLET = ("a:tabLst", "a:defRPr", "a:extLst")
_AUTONUM = {"number": "arabicPeriod", "upper": "alphaUcPeriod", "lower": "alphaLcPeriod"}


def set_bullet(shape, paragraph, kind: str, level: int, start: int = 1) -> None:
    """Give a paragraph the bullet its Markdown asks for: 'p' / 'code' (none),
    'bullet', or a numbered kind. An inherited setting that already matches
    is left alone, so a corporate template's own bullet glyphs stay."""
    want = "none" if kind in ("p", "code") else "char" if kind == "bullet" else kind
    inherited = paragraph_style(shape, paragraph).bullet if shape.is_placeholder else "none"
    if inherited == want and start == 1:
        return
    ppr = paragraph._p.get_or_add_pPr()  # noqa: SLF001
    for tag in _BU_TAGS:
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    if want == "none":
        _insert_bullet(ppr, etree.SubElement(ppr, qn("a:buNone")))
        if not shape.is_placeholder:
            ppr.set("marL", str(_INDENT * level))
            ppr.set("indent", "0")
        return
    if want == "char":
        font = etree.SubElement(ppr, qn("a:buFont"))
        font.set("typeface", "Arial")
        _insert_bullet(ppr, font)
        char = etree.SubElement(ppr, qn("a:buChar"))
        char.set("char", "•")
        _insert_bullet(ppr, char)
    else:
        auto = etree.SubElement(ppr, qn("a:buAutoNum"))
        auto.set("type", _AUTONUM[want])
        if start != 1:
            auto.set("startAt", str(start))
        _insert_bullet(ppr, auto)
    if not shape.is_placeholder or inherited == "none":
        ppr.set("marL", str(_INDENT * (level + 1)))
        ppr.set("indent", str(-_INDENT))


def level_styles(shape) -> list:
    """The inherited style (reader.ParaStyle) of outline levels 0–8 in a
    placeholder. Corporate templates use levels as text styles (plain text,
    subheading, serif line, A. list), so the export picks levels by style."""
    p = shape.text_frame.paragraphs[0]
    styles = []
    for level in range(9):
        p.level = level
        styles.append(paragraph_style(shape, p))
    p.level = 0
    return styles


def text_fill(layout_placeholder):
    """A copy of the text colour (<a:solidFill>) a layout placeholder's first
    level inherits, or None: free text boxes on that layout take it, so they
    stay readable on dark layouts."""
    current = layout_placeholder
    for _ in range(2):  # layout → master placeholder
        if current is None:
            break
        fill = current._element.find(  # noqa: SLF001
            f".//{qn('a:lstStyle')}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}/{qn('a:solidFill')}")
        if fill is not None:
            return copy.deepcopy(fill)
        current = getattr(current, "_base_placeholder", None)
    try:
        master = layout_placeholder.part.slide_layout.slide_master if hasattr(layout_placeholder.part, "slide_layout") \
            else layout_placeholder.part.slide_master
    except AttributeError:
        return None
    fill = master._element.find(  # noqa: SLF001
        f"{qn('p:txStyles')}/{qn('p:bodyStyle')}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}/{qn('a:solidFill')}")
    return copy.deepcopy(fill) if fill is not None else None


def set_run_fill(run, fill) -> None:
    """Give a run a copy of a <a:solidFill> (from text_fill)."""
    rpr = run._r.get_or_add_rPr()  # noqa: SLF001
    for old in rpr.findall(qn("a:solidFill")):
        rpr.remove(old)
    el = copy.deepcopy(fill)
    ln = rpr.find(qn("a:ln"))
    if ln is not None:
        ln.addnext(el)
    else:
        rpr.insert(0, el)


def _insert_bullet(ppr, el) -> None:
    """Move a bullet element before tabLst / defRPr / extLst (schema order)."""
    ppr.remove(el)
    for tag in _AFTER_BULLET:
        later = ppr.find(qn(tag))
        if later is not None:
            later.addprevious(el)
            return
    ppr.append(el)


def insert_paragraph_before(paragraph) -> None:
    """A new, empty paragraph before this one (python-pptx can only append)."""
    p = paragraph._p  # noqa: SLF001
    p.addprevious(p.makeelement(qn("a:p"), {}))


def remove_shape(shape) -> None:
    el = shape._element  # noqa: SLF001
    el.getparent().remove(el)


def set_alt_text(shape, text: str) -> None:
    for cnvpr in shape._element.iter(qn("p:cNvPr")):  # noqa: SLF001
        cnvpr.set("descr", text)
        return


def set_hidden(slide) -> None:
    slide._element.set("show", "0")  # noqa: SLF001


def set_picture_opacity(picture, opacity: float) -> None:
    blip = picture._element.find(".//" + qn("a:blip"))  # noqa: SLF001
    if blip is None:
        return
    mod = etree.Element(qn("a:alphaModFix"))
    mod.set("amt", str(round(max(0.0, min(1.0, opacity)) * 100000)))
    blip.insert(0, mod)


def add_svg(picture, svg: bytes) -> None:
    """Attach the SVG to a picture the way PowerPoint 2016+ does: the raster
    blip stays as the fallback, an asvg:svgBlip extension names the SVG part."""
    slide_part = picture.part
    package = slide_part.package
    partname = package.next_partname("/ppt/media/image%d.svg")
    part = Part(partname, "image/svg+xml", package, svg)
    rid = slide_part.relate_to(part, RT.IMAGE)
    blip = picture._element.find(".//" + qn("a:blip"))  # noqa: SLF001
    ext_lst = blip.find(qn("a:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(blip, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"))
    ext.set("uri", _SVG_URI)
    svg_blip = etree.SubElement(ext, f"{{{_ASVG}}}svgBlip", nsmap={"asvg": _ASVG})
    svg_blip.set(f"{{{_R}}}embed", rid)


def set_fill_alpha(shape, alpha: float) -> None:
    """Fade a shape's solid fill (after fill.solid() + fore_color.rgb)."""
    sppr = shape._element.find(qn("p:spPr"))  # noqa: SLF001
    color = sppr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}") if sppr is not None else None
    if color is not None:
        etree.SubElement(color, qn("a:alpha")).set("val", str(round(max(0.0, min(1.0, alpha)) * 100000)))


def set_text_outline(run, width_pt: float, color: str) -> None:
    """A halo around the letters (text elements over photos)."""
    rpr = run._r.get_or_add_rPr()  # noqa: SLF001
    ln = etree.Element(qn("a:ln"))
    ln.set("w", str(round(width_pt * 12700)))
    fill = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", color.lstrip("#").upper())
    rpr.insert(0, ln)


def set_corner_radius(shape, fraction: float) -> None:
    """Rounded rectangle: adjust value as a fraction of the shorter side (0–0.5)."""
    if len(shape.adjustments):
        shape.adjustments[0] = max(0.0, min(0.5, fraction))


def remove_original_slides(prs, count: int) -> None:
    """Drop the first `count` slides (the base deck's own) after the new ones were added."""
    sld_ids = prs.slides._sldIdLst  # noqa: SLF001
    for sld_id in list(sld_ids)[:count]:
        prs.part.drop_rel(sld_id.get(qn("r:id")))
        sld_ids.remove(sld_id)


def set_sections(prs, sections: list[tuple[str, list[int]]]) -> None:
    """Write PowerPoint sections (p14:sectionLst) from (name, [slide ids])."""
    root = prs.part._element  # noqa: SLF001
    for old in root.findall(f".//{{{P14}}}sectionLst"):
        ext = old.getparent()
        ext.getparent().remove(ext)
    if not sections:
        return
    ext_lst = root.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(root, qn("p:extLst"))
    ext = etree.SubElement(ext_lst, qn("p:ext"))
    ext.set("uri", _SECTIONS_URI)
    lst = etree.SubElement(ext, f"{{{P14}}}sectionLst", nsmap={"p14": P14})
    for name, ids in sections:
        section = etree.SubElement(lst, f"{{{P14}}}section")
        section.set("name", name)
        section.set("id", "{" + str(uuid.uuid4()).upper() + "}")
        id_lst = etree.SubElement(section, f"{{{P14}}}sldIdLst")
        for sid in ids:
            etree.SubElement(id_lst, f"{{{P14}}}sldId").set("id", str(sid))


def clone_shape(shape, slide) -> None:
    """Copy a shape (e.g. a decorative placeholder of a base-deck slide: the
    logo, the slide number) onto another slide, pictures included."""
    el = copy.deepcopy(shape._element)  # noqa: SLF001
    src_part = shape.part
    for node in el.iter():
        for attr in (f"{{{_R}}}embed", f"{{{_R}}}link", f"{{{_R}}}id"):
            rid = node.get(attr)
            if rid and rid in src_part.rels:
                rel = src_part.rels[rid]
                if rel.is_external:
                    node.set(attr, slide.part.relate_to(rel.target_ref, rel.reltype, is_external=True))
                elif rel.reltype == RT.IMAGE:
                    node.set(attr, slide.part.relate_to(rel.target_part, RT.IMAGE))
                else:
                    del node.attrib[attr]
    slide.shapes._spTree.append(el)  # noqa: SLF001


def is_empty_placeholder(shape) -> bool:
    """A placeholder without text and without a picture: shows only its prompt."""
    if shape.has_text_frame:
        return not shape.text_frame.text.strip()
    return shape._element.find(".//" + qn("a:blip")) is None  # noqa: SLF001
