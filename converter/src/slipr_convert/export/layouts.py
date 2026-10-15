"""Which PowerPoint layout a Slipr slide gets, and which placeholder each slot fills.

Template-agnostic, like mapping.py on the import side: it reads the
template's `pptxMapping` backwards and otherwise reasons about slot *types*
and placeholder *types*, never about a template's layout ids.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from pptx.enum.shapes import PP_PLACEHOLDER as PH

TITLE_TYPES = {PH.TITLE, PH.CENTER_TITLE, PH.VERTICAL_TITLE}
CHROME_TYPES = {PH.DATE, PH.FOOTER, PH.SLIDE_NUMBER, PH.HEADER}
_CARD_ROLE = re.compile(r"^(\w+)\[(\d+)\](?:\.(\w+))?$")

# Placeholder types each kind of slot may go into, best first.
_TARGETS = {
    "title": (PH.TITLE, PH.CENTER_TITLE, PH.VERTICAL_TITLE),
    "text": (PH.SUBTITLE, PH.BODY),
    "body": (PH.OBJECT, PH.BODY),
    "image": (PH.PICTURE, PH.OBJECT),
    "visual": (PH.OBJECT,),
}
# Short text slots claim a placeholder in this order; the rest follow in template order.
_TEXT_PRIORITY = ("subtitle", "byline", "caption")
_STANDARD_FALLBACKS = ("Title Only", "Title and Content", "Blank")


@dataclass
class Plan:
    layout: object  # pptx SlideLayout
    reason: str
    roles: dict[int, str | None] = field(default_factory=dict)  # placeholder idx → role, from the template


def slot_kind(slot_type: str, name: str) -> str:
    if slot_type == "text":
        return "title" if name == "title" else "text"
    if slot_type in ("list", "richtext", "cards", "quote", "code"):
        return "body"
    if slot_type == "image":
        return "image"
    return "visual"


def choose_layout(slipr_layout: str, prs, template: dict, *, slots: list[tuple[str, str]], first: bool) -> Plan:
    """The PowerPoint layout for a Slipr layout: of the `pptxMapping` entries
    that map to it and exist in the deck, the one whose placeholders take the
    most of the filled slots (ties: the first). Else a layout chosen by shape.

    slots: (name, type) of the slots that have content, in template order.
    """
    by_name = {layout.name: layout for layout in prs.slide_layouts}
    ignore = set(template.get("pptxIgnore") or [])
    best: tuple[int, Plan] | None = None
    for name, value in (template.get("pptxMapping") or {}).items():
        target = value if isinstance(value, str) else value.get("layout")
        if target != slipr_layout or name not in by_name:
            continue
        roles = {int(k): v for k, v in ((value.get("placeholders") or {}) if isinstance(value, dict) else {}).items()}
        plan = Plan(by_name[name], "template pptxMapping", roles)
        missing = _unplaced(plan, slots, ignore)
        if best is None or missing < best[0]:
            best = (missing, plan)
    if best is not None:
        return best[1]
    has_body = any(slot_kind(t, n) not in ("title", "text") for n, t in slots)
    if not has_body:
        wanted = "Title Slide" if first else "Section Header"
        if wanted in by_name:
            return Plan(by_name[wanted], "title-only slide")
    for name in _STANDARD_FALLBACKS:
        if name in by_name:
            return Plan(by_name[name], "no mapping: content drawn under the title")
    # A deck with its own layout names: the one with a title and the fewest other placeholders.
    def cost(layout) -> tuple[int, int]:
        types = [p.placeholder_format.type for p in layout.placeholders if p.placeholder_format.type not in CHROME_TYPES]
        has_title = any(t in TITLE_TYPES for t in types)
        return (0 if has_title else 1, len(types))
    layout = min(prs.slide_layouts, key=cost)
    return Plan(layout, "no mapping: the simplest layout with a title")


def _unplaced(plan: Plan, slots: list[tuple[str, str]], ignore: set[int]) -> int:
    """How many filled slots would get no placeholder on this layout."""
    by_role = {parse_role(r)[0] for r in plan.roles.values() if r}
    taken = set(plan.roles) | ignore
    rest = [(n, t) for n, t in slots if n not in by_role]
    free = [p for p in plan.layout.placeholders if p.placeholder_format.type not in CHROME_TYPES]
    return len(rest) - len(assign(rest, free, taken))


def parse_role(role: str) -> tuple[str, int | None, str | None]:
    """'title' → ('title', None, None); 'cards[2].name' → ('cards', 2, 'name')."""
    m = _CARD_ROLE.match(role)
    if m:
        return m.group(1), int(m.group(2)), m.group(3)
    return role, None, None


def assign(slots: list[tuple[str, str]], placeholders: list, taken: set[int]) -> dict[str, object]:
    """Slot name → placeholder for the slots without a template role.

    slots: (name, type) of the slots that have content, in template order.
    placeholders: the new slide's free placeholders (not chrome, not ignored).
    """
    free = [p for p in placeholders if p.placeholder_format.idx not in taken]
    free.sort(key=lambda p: p.placeholder_format.idx)
    kinds = [(n, t, slot_kind(t, n)) for n, t in slots]
    titles = [s for s in kinds if s[2] == "title"]
    texts = [s for s in kinds if s[2] == "text"]
    texts.sort(key=lambda s: (_TEXT_PRIORITY.index(s[0]) if s[0] in _TEXT_PRIORITY else len(_TEXT_PRIORITY),
                              1 if s[0] == "kicker" else 0))
    others = [s for s in kinds if s[2] in ("body", "image", "visual")]
    result: dict[str, object] = {}
    for name, _type, kind in titles + others + texts:
        for ph_type in _TARGETS[kind]:
            match = next((p for p in free if p.placeholder_format.type == ph_type), None)
            if match is not None:
                result[name] = match
                free.remove(match)
                break
    return result
