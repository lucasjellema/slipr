"""The export core: a Slipr show in, a .pptx (bytes) and a report out. No files written.

    from slipr_convert.export import export_show
    result = export_show(show, template, read_file=lambda rel: (folder / rel).read_bytes())
    result.pptx      # bytes of the deck
    result.report    # what went where, and what PowerPoint could not take
"""
from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from . import data as datafiles
from . import markdown as md
from . import xml
from .layouts import CHROME_TYPES, TITLE_TYPES, assign, choose_layout, parse_role, slot_kind
from .report import ExportReport, ExportSlideReport

ReadFile = datafiles.ReadFile

_SIZES = {"16:9": (Emu(12192000), Emu(6858000)), "4:3": (Emu(9144000), Emu(6858000))}
# The template sizes of text elements, as a percentage of the slide width (the default theme's --slipr-size-*).
_NAMED_SIZES = {"small": 1.2, "body": 2.2, "subtitle": 2.1, "title": 3.6, "hero": 5.2}
_FONTS = {"heading": "+mj-lt", "body": "+mn-lt", "mono": "Consolas", "serif": "Georgia", "system": "Segoe UI"}
_RASTER = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"}
_CONVERTIBLE = {"webp"}  # turned into PNG with Pillow (a python-pptx dependency)
_VIDEO_MIME = {"mp4": "video/mp4", "m4v": "video/mp4", "webm": "video/webm", "mov": "video/quicktime"}
_MAX_TABLE_ROWS = 16
_KIND_LABEL = {"diagram": "Diagram", "graph": "Graph", "map": "Map", "web": "Web page", "pdf": "PDF document",
               "video": "Video", "table": "Table", "tag-cloud": "Tag cloud", "chart": "Chart"}


@dataclass
class ExportResult:
    pptx: bytes
    report: ExportReport


@dataclass
class Box:
    left: int
    top: int
    width: int
    height: int

    @classmethod
    def of(cls, shape) -> "Box":
        return cls(int(shape.left or 0), int(shape.top or 0), int(shape.width or 0), int(shape.height or 0))

    def rows(self, n: int, gap: int) -> list["Box"]:
        h = (self.height - gap * (n - 1)) // max(1, n)
        return [Box(self.left, self.top + i * (h + gap), self.width, h) for i in range(n)]

    def columns(self, fractions: list[float], gap: int) -> list["Box"]:
        usable = self.width - gap * (len(fractions) - 1)
        out, x = [], self.left
        for f in fractions:
            w = int(usable * f)
            out.append(Box(x, self.top, w, self.height))
            x += w + gap
        return out


def export_show(
    show: dict,
    template: dict,
    *,
    read_file: ReadFile,
    base: bytes | None = None,
    include_notes: bool = True,
    include_reader: bool = True,
    source_name: str = "show.json",
) -> ExportResult:
    """Build a PowerPoint deck from a show.

    base: a .pptx whose layouts (and masters, theme) the deck uses, such as the
    company template a Slipr template was derived from. Its own slides are
    dropped; the first slide of each layout lends its decoration (logo, slide
    number) to new slides on that layout. Without it, python-pptx's default
    Office layouts are used, at the show's aspect ratio.
    """
    report = ExportReport(source_name)
    prs = Presentation(io.BytesIO(base)) if base else Presentation()
    if not base:
        prs.slide_width, prs.slide_height = _SIZES.get((show.get("settings") or {}).get("aspectRatio", "16:9"), _SIZES["16:9"])
    samples: dict[str, object] = {}
    for s in prs.slides:
        samples.setdefault(s.slide_layout.name, s)
    original_count = len(prs.slides)
    ctx = _Context(prs, show, template, read_file, include_notes, include_reader)

    # Chapters become PowerPoint sections. A slide without a chapter stays in the
    # section before it (PowerPoint puts every slide in a section once there are any).
    chapters = {c.get("id"): c.get("title") or c.get("id") for c in show.get("chapters") or []}
    sections: list[tuple[str, list[int]]] = []
    current = None
    for number, slide in enumerate(show.get("slides") or [], start=1):
        entry = report.slide(number, slide.get("id", f"s-{number}"), str((slide.get("content") or {}).get("title") or ""))
        pptx_slide = ctx.add_slide(slide, number == 1, samples, entry)
        chapter = slide.get("chapter")
        if chapter in chapters and chapter != current:
            current = chapter
            sections.append((chapters[chapter], []))
        elif not sections:
            sections.append(("Introduction", []))
        sections[-1][1].append(pptx_slide.slide_id)

    if base:
        xml.remove_original_slides(prs, original_count)
    xml.set_sections(prs, sections if any(name != "Introduction" for name, _ in sections) else [])

    meta = show.get("meta") or {}
    props = prs.core_properties
    props.title = meta.get("title") or ""
    props.subject = meta.get("subtitle") or meta.get("description") or ""
    props.author = meta.get("author") or ""
    props.keywords = ", ".join(meta.get("tags") or [])
    props.language = meta.get("language") or ""
    props.comments = "Exported from the Slipr show " + str(show.get("id", ""))

    out = io.BytesIO()
    prs.save(out)
    return ExportResult(out.getvalue(), report)


class _Context:
    def __init__(self, prs, show: dict, template: dict, read_file: ReadFile, include_notes: bool, include_reader: bool):
        self.prs = prs
        self.show = show
        self.template = template
        self.read_file = read_file
        self.include_notes = include_notes
        self.include_reader = include_reader
        self.glossary = show.get("glossary") or {}
        self.layouts = {layout["id"]: layout for layout in template.get("layouts") or []}
        self.ignore = set(template.get("pptxIgnore") or [])
        self.W = int(prs.slide_width)
        self.H = int(prs.slide_height)
        self.gap = int(self.W * 0.02)
        self.notes_extra: list[str] = []
        self.text_fill = None
        self._svg_noted = None

    # -- one slide ---------------------------------------------------------------------------------

    def add_slide(self, slide: dict, first: bool, samples: dict, entry: ExportSlideReport):
        content = slide.get("content") or {}
        layout_def = self.layouts.get(slide.get("layout"), {"slots": []})
        slot_types = {s["name"]: s for s in layout_def.get("slots") or []}
        filled = [(name, slot_types[name]["type"]) for name in slot_types if _has_value(content.get(name))]
        for name in content:
            if name not in slot_types and _has_value(content[name]):
                entry.notes.append(f"slot `{name}` is not in layout `{slide.get('layout')}`, exported as text")
                filled.append((name, "richtext" if isinstance(content[name], str) else "unknown"))
        plan = choose_layout(slide.get("layout", ""), self.prs, self.template, slots=filled, first=first)
        entry.layout, entry.reason = plan.layout.name, plan.reason
        pslide = self.prs.slides.add_slide(plan.layout)
        self.notes_extra = []

        ignored = {idx for idx, role in plan.roles.items() if role is None} | (self.ignore - {i for i, r in plan.roles.items() if r})
        # Free text boxes take the text colour of the layout's body text (white on dark layouts).
        reference = sorted((p for p in plan.layout.placeholders if p.placeholder_format.idx not in ignored
                            and p.placeholder_format.type not in CHROME_TYPES),
                           key=lambda p: (p.placeholder_format.type in TITLE_TYPES, p.placeholder_format.idx))
        self.text_fill = xml.text_fill(reference[0]) if reference else None
        self._decorate(pslide, samples.get(plan.layout.name), ignored)
        placeholders = {p.placeholder_format.idx: p for p in pslide.placeholders}
        used: set[int] = set()
        done: set[str] = set()

        # 1. Template roles (idx → slot), as the converter reads them.
        cards_by_index: dict[str, dict[int, list]] = {}
        for idx, role in plan.roles.items():
            if not role or idx not in placeholders:
                continue
            slot, card, field = parse_role(role)
            if slot not in content or not _has_value(content[slot]):
                continue
            if card is None:
                self._fill(pslide, placeholders[idx], slot, slot_types.get(slot, {"type": "richtext"}), content[slot], entry)
                done.add(slot)
            else:
                cards_by_index.setdefault(slot, {}).setdefault(card, []).append((idx, field))
            used.add(idx)
        for slot, by_card in cards_by_index.items():
            cards = content[slot] if isinstance(content[slot], list) else []
            fields = (slot_types.get(slot) or {}).get("fields") or []
            for card_index, targets in by_card.items():
                card = cards[card_index] if card_index < len(cards) else {}
                for idx, field in targets:
                    self._fill_card_placeholder(pslide, placeholders[idx], card, fields, field, entry)
            extra = [c for i, c in enumerate(cards) if i not in by_card]
            if extra:
                entry.notes.append(f"{len(extra)} card(s) of `{slot}` have no placeholder, exported under the slide")
                self._cards_grid(pslide, extra, fields, self._free_area(pslide).rows(1, 0)[0], entry)
            done.add(slot)

        # Filling a picture placeholder replaces its element: look the placeholders up again.
        placeholders = {p.placeholder_format.idx: p for p in pslide.placeholders}

        # 2. The off-slide title some corporate templates have (role null): keep the outline title.
        title_text = content.get("title")
        for idx, p in placeholders.items():
            if p.placeholder_format.type in TITLE_TYPES and idx not in used and plan.roles.get(idx, "") is None:
                if title_text and self._off_slide(p):
                    p.text_frame.text = str(title_text)
                    used.add(idx)

        # 3. Generic: the remaining slots by type into free placeholders.
        candidates = [p for idx, p in placeholders.items()
                      if idx not in used and idx not in ignored and p.placeholder_format.type not in CHROME_TYPES
                      and not self._off_slide(p)]
        rest = [(n, t) for n, t in filled if n not in done]
        target_idx = {name: ph.placeholder_format.idx for name, ph in assign(rest, candidates, used).items()}
        for name, idx in target_idx.items():
            self._fill(pslide, placeholders[idx], name, slot_types.get(name, {"type": "richtext"}), content[name], entry)
            used.add(idx)
            done.add(name)
        current = {p.placeholder_format.idx: p for p in pslide.placeholders}
        targets = {name: current[idx] for name, idx in target_idx.items() if idx in current}

        # 4. Short texts without a placeholder join a text placeholder; everything else goes under the title.
        leftovers = [(n, t) for n, t in filled if n not in done]
        free_items = []
        for name, slot_type in leftovers:
            host = self._text_host(targets, slot_types)
            paired = targets.get(name[:-5]) if name.endswith("Title") else None  # leftTitle → left
            if slot_kind(slot_type, name) == "text" and paired is not None and paired.has_text_frame:
                self._prepend_heading(paired, content[name])
                entry.notes.append(f"`{name}` became the heading of `{name[:-5]}`")
            elif slot_kind(slot_type, name) == "text" and host is not None:
                self._merge_text(host, content[name], before=name == "kicker")
                entry.notes.append(f"`{name}` joined the text of `{host.name}`")
            else:
                free_items.append((name, slot_type))
        if free_items:
            self._place_free(pslide, free_items, content, slot_types, entry)

        # 5. Placeholders nobody filled show "Click to add text" in PowerPoint: remove them.
        for p in list(pslide.placeholders):
            idx = p.placeholder_format.idx
            if idx not in used and idx not in ignored and p.placeholder_format.type not in CHROME_TYPES and xml.is_empty_placeholder(p):
                try:
                    xml.remove_shape(p)
                except ValueError:
                    pass

        for overlay in slide.get("overlays") or []:
            self._overlay(pslide, overlay, entry)
        if (slide.get("audience") or {}).get("hidden"):
            xml.set_hidden(pslide)
            entry.notes.append("hidden (backup slide)")
        self._notes(pslide, slide)
        return pslide

    def _decorate(self, pslide, sample, ignored: set[int]) -> None:
        """Copy the base deck's decoration (logo, slide number, filled decorative placeholders)."""
        if sample is None:
            return
        mine = {p.placeholder_format.idx: p for p in pslide.placeholders}
        for p in sample.placeholders:
            idx = p.placeholder_format.idx
            kind = p.placeholder_format.type
            if (idx in ignored or kind in (PH.SLIDE_NUMBER, PH.FOOTER)) and kind not in TITLE_TYPES and not xml.is_empty_placeholder(p):
                if idx in mine:
                    xml.remove_shape(mine[idx])
                xml.clone_shape(p, pslide)

    def _off_slide(self, shape) -> bool:
        b = Box.of(shape)
        return b.left + b.width <= 0 or b.top + b.height <= 0 or b.left >= self.W or b.top >= self.H

    def _text_host(self, targets: dict, slot_types: dict):
        for name, ph in targets.items():
            if slot_types.get(name, {}).get("type") == "text" and name != "title" and ph.has_text_frame:
                return ph
        return None

    def _prepend_heading(self, shape, value) -> None:
        paras = [md.Para("p", 0, [md.Run(str(value), bold=True)])]
        existing = shape.text_frame.paragraphs
        xml.insert_paragraph_before(existing[0])
        self._write_runs(shape, shape.text_frame.paragraphs[0], paras[0])

    def _merge_text(self, shape, value, *, before: bool) -> None:
        lines = [p.text for p in shape.text_frame.paragraphs if p.text]
        new = str(value)
        lines = [new, *lines] if before else [*lines, new]
        self._write_paras(shape, [md.Para("p", 0, [md.Run(line)]) for line in lines])

    # -- slot values --------------------------------------------------------------------------------

    def _fill(self, pslide, ph, name: str, slot: dict, value, entry: ExportSlideReport) -> None:
        slot_type = slot.get("type")
        kind = slot_kind(slot_type, name)
        if kind == "image":
            self._image(pslide, value, Box.of(ph), entry, placeholder=ph, label=name)
        elif kind == "visual":
            box = Box.of(ph)
            xml.remove_shape(ph)
            self._visual(pslide, slot_type, value, box, entry, label=name)
        elif slot_type == "cards":
            box = Box.of(ph)
            xml.remove_shape(ph)
            self._cards_grid(pslide, value if isinstance(value, list) else [], slot.get("fields") or [], box, entry)
        elif ph.has_text_frame:
            self._write_paras(ph, self._paras(slot_type, value))
        else:
            box = Box.of(ph)
            xml.remove_shape(ph)
            self._textbox(pslide, box, self._paras(slot_type, value))

    def _paras(self, slot_type: str | None, value) -> list[md.Para]:
        if slot_type == "text":
            return [md.Para("p", 0, [md.Run(value)], role="line")] if isinstance(value, str) else []
        if slot_type == "quote" and isinstance(value, dict):
            paras = [md.Para("p", 0, [md.Run("“" + str(value.get("text", "")) + "”", italic=True)])]
            if value.get("source"):
                paras.append(md.Para("p", 0, [md.Run("— " + value["source"])]))
            return paras
        if slot_type == "code" and isinstance(value, dict):
            return [md.Para("code", 0, [md.Run(line, code=True)]) for line in str(value.get("code", "")).split("\n")]
        if isinstance(value, (str, list)):
            return md.parse(value, self.glossary)
        return [md.Para("p", 0, [md.Run(str(value))])]

    def _fill_card_placeholder(self, pslide, ph, card: dict, fields: list[dict], field: str | None, entry) -> None:
        types = {f["name"]: f.get("type", "text") for f in fields}
        if field is not None:
            value = card.get(field)
            if not _has_value(value):
                return
            if types.get(field) == "image":
                self._image(pslide, value, Box.of(ph), entry, placeholder=ph, label=f"card {field}")
            elif ph.has_text_frame:
                self._write_paras(ph, self._paras(types.get(field, "richtext"), value))
            return
        if ph.placeholder_format.type == PH.PICTURE:
            image = next((card.get(f["name"]) for f in fields if f.get("type") == "image" and _has_value(card.get(f["name"]))), None)
            if image:
                self._image(pslide, image, Box.of(ph), entry, placeholder=ph, label="card picture")
            return
        paras: list[md.Para] = []
        for f in fields:
            if f.get("type") != "image" and _has_value(card.get(f["name"])):
                part = self._paras(f.get("type", "text"), card[f["name"]])
                first = not paras
                for p in part:  # the first field takes the placeholder's first level, the rest its text style
                    p.role = "heading" if first else ""
                paras += part
        if paras and ph.has_text_frame:
            self._write_paras(ph, paras)

    def _cards_grid(self, pslide, cards: list, fields: list[dict], box: Box, entry) -> None:
        cards = [c for c in cards if isinstance(c, dict) and any(_has_value(v) for v in c.values())]
        if not cards:
            return
        n = len(cards)
        cols = n if n <= 4 else 3 if n in (5, 6, 9) else 4
        rows = math.ceil(n / cols)
        cells = [c for row in box.rows(rows, self.gap) for c in row.columns([1 / cols] * cols, self.gap)]
        image_field = next((f["name"] for f in fields if f.get("type") == "image"), None)
        text_fields = [f for f in fields if f.get("type") != "image"] or [{"name": k, "type": "text"} for k in cards[0]]
        size = Pt(16 if rows == 1 else 13)
        for card, cell in zip(cards, cells):
            text_box = cell
            if image_field and _has_value(card.get(image_field)):
                pic_box, text_box = Box(cell.left, cell.top, cell.width, int(cell.height * 0.45)), \
                    Box(cell.left, cell.top + int(cell.height * 0.45) + self.gap // 2, cell.width, int(cell.height * 0.55) - self.gap // 2)
                self._image(pslide, {**card[image_field], "fit": card[image_field].get("fit", "cover")}, pic_box, entry, label="card picture")
            paras: list[md.Para] = []
            for i, f in enumerate(text_fields):
                value = card.get(f["name"])
                if not _has_value(value):
                    continue
                part = self._paras(f.get("type", "text"), value)
                if i == 0 and f.get("type", "text") == "text":
                    for r in (run for p in part for run in p.runs):
                        r.bold = True
                paras += part
            self._textbox(pslide, text_box, paras, size=size)

    # -- text ---------------------------------------------------------------------------------------

    def _write_paras(self, shape, paras: list[md.Para], *, size=None, font=None) -> None:
        tf = shape.text_frame
        tf.clear()
        if not paras:
            return
        styles = xml.level_styles(shape) if shape.is_placeholder else None
        in_list: set[int] = set()  # levels with a numbered list running: PowerPoint continues its numbers
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            first_number = para.kind in ("number", "upper", "lower") and para.level not in in_list
            level = _pick_level(styles, para) if styles else None
            if para.role == "heading" and level != 0:  # the layout has no heading level: bold instead
                for r in para.runs:
                    r.bold = True
            self._write_runs(shape, p, para, size=size, font=font, start=para.start if first_number else 1, level=level)
            if para.kind in ("number", "upper", "lower"):
                in_list.add(para.level)
            else:
                in_list.discard(para.level)

    def _write_runs(self, shape, p, para: md.Para, *, size=None, font=None, start: int = 1, level: int | None = None) -> None:
        """One paragraph. `level` is the outline level whose inherited style
        matches the paragraph (see _pick_level); without one, the Markdown
        level is used and the bullet is set on the paragraph itself."""
        if level is not None:
            p.level = level
        elif para.kind not in ("p", "code"):
            p.level = min(para.level, 8)
        for run in para.runs:
            if run.is_break:
                p.add_line_break()
                continue
            r = p.add_run()
            r.text = run.text
            if run.bold:
                r.font.bold = True
            if run.italic:
                r.font.italic = True
            if run.code or para.kind == "code":
                r.font.name = "Consolas"
            elif font:
                r.font.name = font
            if size is not None:
                r.font.size = size
            if run.link:
                r.hyperlink.address = run.link
            elif not shape.is_placeholder and self.text_fill is not None:
                xml.set_run_fill(r, self.text_fill)
        xml.set_bullet(shape, p, para.kind, p.level if level is not None else para.level, start)

    def _textbox(self, pslide, box: Box, paras: list[md.Para], *, size=None, align=None, anchor=None):
        shape = pslide.shapes.add_textbox(box.left, box.top, box.width, box.height)
        tf = shape.text_frame
        tf.word_wrap = True
        if anchor is not None:
            tf.vertical_anchor = anchor
        self._write_paras(shape, paras, size=size or Pt(18))
        if align is not None:
            for p in tf.paragraphs:
                p.alignment = align
        return shape

    def _card(self, pslide, box: Box, title: str, lines: list[md.Para], link: str | None = None):
        """A link card for content PowerPoint cannot show: a framed box with a title and a few lines."""
        shape = pslide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box.left, box.top, box.width, box.height)
        shape.fill.background()
        shape.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        shape.line.width = Pt(1.5)
        xml.set_corner_radius(shape, 0.06)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        paras = [md.Para("p", 0, [md.Run(title, bold=True, link=link)])] + lines
        self._write_paras(shape, paras, size=Pt(18))
        for n, p in enumerate(tf.paragraphs):
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                if n == 0:
                    r.font.size = Pt(24)
                if self.text_fill is None:
                    r.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
        if link:
            shape.click_action.hyperlink.address = link
        return shape

    # -- pictures -----------------------------------------------------------------------------------

    def _image_blob(self, value: dict, entry, label: str) -> tuple[bytes, bytes | None] | None:
        """(raster picture, SVG or None). An SVG gets a plain PNG stand-in of its
        size: PowerPoint 2016 and later show the SVG, older apps the stand-in."""
        src = (value or {}).get("src") or ""
        ext = src.rsplit(".", 1)[-1].lower() if "." in src else ""
        if re.match(r"^(https?:|data:)", src, re.I):
            entry.skipped.append(f"{label}: picture at an address ({src[:60]}), not a file in the show")
            return None
        blob = self.read_file(src) if src else None
        if blob is None:
            entry.skipped.append(f"{label}: picture not found ({src})")
            return None
        if ext == "svg":
            w, h = _svg_size(blob)
            from PIL import Image
            out = io.BytesIO()
            scale = 400 / max(w, h)
            Image.new("RGB", (max(1, round(w * scale)), max(1, round(h * scale))), (224, 224, 224)).save(out, "PNG")
            if self._svg_noted is not entry:
                entry.notes.append("SVG pictures kept as SVG (PowerPoint 2016 and later; older apps show a grey box)")
                self._svg_noted = entry
            return out.getvalue(), blob
        if ext in _CONVERTIBLE:
            try:
                from PIL import Image
                out = io.BytesIO()
                Image.open(io.BytesIO(blob)).save(out, "PNG")
                return out.getvalue(), None
            except Exception as err:  # Pillow without WebP support, or a broken file
                entry.skipped.append(f"{label}: cannot convert {src} to PNG ({err})")
                return None
        if ext not in _RASTER:
            entry.skipped.append(f"{label}: {ext.upper() or 'this'} pictures cannot go into PowerPoint ({src}); add a PNG version")
            return None
        return blob, None

    def _image(self, pslide, value, box: Box, entry, *, placeholder=None, label: str = "picture"):
        if not isinstance(value, dict):
            return None
        blobs = self._image_blob(value, entry, label)
        blob, svg = blobs if blobs else (None, None)
        if blob is None:
            if placeholder is not None and placeholder.placeholder_format.type != PH.PICTURE:
                xml.remove_shape(placeholder)
            self._card(pslide, box, value.get("alt") or "Picture", [md.Para("p", 0, [md.Run(value.get("src", ""))])])
            return None
        from PIL import Image
        with Image.open(io.BytesIO(blob)) as im:
            iw, ih = im.size
        crop = value.get("crop") or {"x": 0, "y": 0, "w": 1, "h": 1}
        cx, cy = float(crop.get("x", 0)), float(crop.get("y", 0))
        cw, ch = max(0.01, float(crop.get("w", 1))), max(0.01, float(crop.get("h", 1)))
        visible_w, visible_h = iw * cw, ih * ch
        in_picture_ph = placeholder is not None and placeholder.placeholder_format.type == PH.PICTURE
        fit = value.get("fit") or ("cover" if in_picture_ph else "contain")
        if fit == "cover":
            box_ratio = box.width / max(1, box.height)
            img_ratio = visible_w / max(1, visible_h)
            fx, fy = _focus(value.get("focus"))
            if img_ratio > box_ratio:  # wider than the box: crop left and right
                keep = cw * box_ratio / img_ratio
                cx, cw = cx + (cw - keep) * fx, keep
            else:
                keep = ch * img_ratio / box_ratio
                cy, ch = cy + (ch - keep) * fy, keep
            if in_picture_ph:
                pic = placeholder.insert_picture(io.BytesIO(blob))
            else:
                if placeholder is not None:
                    xml.remove_shape(placeholder)
                pic = pslide.shapes.add_picture(io.BytesIO(blob), box.left, box.top, box.width, box.height)
        else:
            if placeholder is not None:
                xml.remove_shape(placeholder)
            scale = min(box.width / visible_w, box.height / visible_h)
            w, h = int(visible_w * scale), int(visible_h * scale)
            pic = pslide.shapes.add_picture(io.BytesIO(blob), box.left + (box.width - w) // 2, box.top + (box.height - h) // 2, w, h)
        pic.crop_left, pic.crop_top = cx, cy
        pic.crop_right, pic.crop_bottom = max(0.0, 1 - cx - cw), max(0.0, 1 - cy - ch)
        if svg is not None:
            xml.add_svg(pic, svg)
        if value.get("alt"):
            xml.set_alt_text(pic, value["alt"])
        if isinstance(value.get("opacity"), (int, float)) and value["opacity"] < 1:
            xml.set_picture_opacity(pic, float(value["opacity"]))
        return pic

    # -- rich slots ---------------------------------------------------------------------------------

    def _visual(self, pslide, slot_type: str, value, box: Box, entry, *, label: str) -> None:
        if not isinstance(value, dict):
            return
        kind_label = _KIND_LABEL.get(slot_type, slot_type)
        title = value.get("title") or value.get("caption") or kind_label
        if slot_type == "table" and not value.get("fallback"):
            self._table(pslide, value, box, entry)
            return
        if slot_type == "tag-cloud" and not value.get("fallback"):
            self._cloud(pslide, value, box, entry)
            return
        if slot_type == "video" and self._video(pslide, value, box, entry):
            return
        url = value.get("url") or value.get("video") or value.get("pdf")
        link = url if isinstance(url, str) and re.match(r"^https?:", url, re.I) else None
        if value.get("fallback"):
            pic = self._image(pslide, value["fallback"], box, entry, label=f"{label} fallback")
            if pic is not None:
                if link:
                    pic.click_action.hyperlink.address = link
                entry.fallbacks.append(f"{kind_label} `{label}` as its fallback picture" + (" (click opens the link)" if link else ""))
                return
        lines: list[md.Para] = []
        if slot_type == "diagram":
            lines.append(md.Para("p", 0, [md.Run("A Mermaid diagram, drawn in the Slipr show. Its source is in the notes.", italic=True)]))
            self.notes_extra.append("Diagram source (Mermaid):\n" + str(value.get("mermaid", "")))
        elif slot_type == "graph":
            self._graph_list(pslide, value, box, entry, label)
            return
        elif slot_type == "map":
            self._map_list(pslide, value, box, entry, label)
            return
        elif slot_type == "pdf":
            page = value.get("pages") or value.get("page")
            where = str(value.get("pdf", "")) + (f", page {page[0]}–{page[1]}" if isinstance(page, list) else f", page {page}" if page else "")
            lines.append(md.Para("p", 0, [md.Run(where)]))
        elif url:
            lines.append(md.Para("p", 0, [md.Run(str(url), link=link)]))
        self._card(pslide, box, title, lines, link)
        entry.fallbacks.append(f"{kind_label} `{label}` as a card" + (" with its link" if link else ""))

    def _table(self, pslide, value: dict, box: Box, entry) -> None:
        labels, rows, totals, align, problem = datafiles.table_cells(value, self.read_file)
        if problem or not labels:
            entry.skipped.append(f"table: {problem or 'no columns'}")
            self._card(pslide, box, value.get("title") or "Table", [md.Para("p", 0, [md.Run(problem or "")])])
            return
        if len(rows) > _MAX_TABLE_ROWS:
            entry.fallbacks.append(f"table: the first {_MAX_TABLE_ROWS} of {len(rows)} rows (all rows are in the show)")
            rows = rows[:_MAX_TABLE_ROWS]
        body = [labels] + rows + ([totals] if totals else [])
        size = Pt(14 if len(body) <= 8 else 12 if len(body) <= 12 else 10)
        row_h = min(box.height // len(body), int(Pt(size.pt * 2.2)))
        frame = pslide.shapes.add_table(len(body), len(labels), box.left, box.top, box.width, row_h * len(body))
        table = frame.table
        for r, cells in enumerate(body):
            table.rows[r].height = row_h
            for c, text in enumerate(cells):
                cell = table.cell(r, c)
                cell.text = str(text)
                for p in cell.text_frame.paragraphs:
                    p.alignment = {"r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}.get(align[c], PP_ALIGN.LEFT) if r else PP_ALIGN.LEFT
                    for run in p.runs:
                        run.font.size = size
                        if totals and r == len(body) - 1:
                            run.font.bold = True
        if value.get("caption"):
            self.notes_extra.append("Table caption: " + str(value["caption"]))

    def _cloud(self, pslide, value: dict, box: Box, entry) -> None:
        terms, problem = datafiles.cloud_terms(value, self.read_file)
        if problem or not terms:
            entry.skipped.append(f"tag cloud: {problem or 'no terms'}")
            return
        weights = [w for _, w in terms]
        low, high = min(weights), max(weights)
        runs: list[md.Run] = []
        for i, (text, _w) in enumerate(terms):
            if i:
                runs.append(md.Run("   "))
            runs.append(md.Run(text))
        shape = self._textbox(pslide, box, [md.Para("p", 0, runs)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        term_runs = [r for r in shape.text_frame.paragraphs[0].runs if r.text.strip()]
        for run, (_t, w) in zip(term_runs, terms):
            t = 0.5 if high == low else (w - low) / (high - low)
            run.font.size = Pt(round(14 + 26 * t))
            run.font.bold = t > 0.66
            run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1 if t > 0.5 else MSO_THEME_COLOR.TEXT_1

    def _graph_list(self, pslide, value: dict, box: Box, entry, label: str) -> None:
        nodes = list(value.get("nodes") or [])
        edges = list(value.get("edges") or [])
        if value.get("data"):
            blob = self.read_file(value["data"])
            if blob:
                import json
                try:
                    extra = json.loads(blob.decode("utf-8"))
                    nodes, edges = list(extra.get("nodes") or []) + nodes, list(extra.get("edges") or []) + edges
                except (ValueError, AttributeError):
                    entry.skipped.append(f"graph: {value['data']} is not valid JSON")
        names = {n.get("id"): n.get("label") or n.get("id") for n in nodes}
        paras = []
        for n in nodes:
            links = [names.get(e.get("to"), e.get("to")) for e in edges if e.get("from") == n.get("id")]
            text = str(names[n.get("id")]) + (" → " + ", ".join(str(x) for x in links) if links else "")
            paras.append(md.Para("bullet", 0, [md.Run(text)]))
        self._textbox(pslide, box, paras, size=Pt(14 if len(paras) > 8 else 18))
        entry.fallbacks.append(f"graph `{label}` as a list of its nodes and links")

    def _map_list(self, pslide, value: dict, box: Box, entry, label: str) -> None:
        paras = [md.Para("bullet", 0, [md.Run(f"{m.get('title') or m.get('id') or 'Place'} ({m.get('lat')}, {m.get('lng')})")])
                 for m in value.get("markers") or []]
        paras += [md.Para("bullet", 0, [md.Run(f"Layer: {layer.get('title') or layer.get('data')}")]) for layer in value.get("layers") or []]
        if not paras and value.get("center"):
            c = value["center"]
            paras = [md.Para("p", 0, [md.Run(f"Map around {c.get('lat')}, {c.get('lng')}")])]
        self._textbox(pslide, box, paras, size=Pt(14 if len(paras) > 8 else 18))
        entry.fallbacks.append(f"map `{label}` as a list of its places")

    def _video(self, pslide, value: dict, box: Box, entry) -> bool:
        """A video file in the show folder becomes a real PowerPoint video."""
        src = str(value.get("video") or "")
        ext = src.rsplit(".", 1)[-1].lower() if "." in src else ""
        if re.match(r"^https?:", src, re.I) or ext not in _VIDEO_MIME:
            return False
        blob = self.read_file(src)
        if blob is None:
            entry.skipped.append(f"video not found ({src})")
            return False
        poster = None
        if value.get("fallback"):
            poster_blobs = self._image_blob(value["fallback"], entry, "video poster")
            poster = io.BytesIO(poster_blobs[0]) if poster_blobs else None
        w, h = box.width, min(box.height, box.width * 9 // 16)
        pslide.shapes.add_movie(io.BytesIO(blob), box.left, box.top + (box.height - h) // 2, w, h,
                                poster_frame_image=poster, mime_type=_VIDEO_MIME[ext])
        if value.get("start") or value.get("end"):
            entry.notes.append("video start / end times are not set in PowerPoint; trim the video there")
        entry.notes.append(f"video {src} embedded")
        return True

    # -- free placement -----------------------------------------------------------------------------

    def _free_area(self, pslide) -> Box:
        title = next((p for p in pslide.placeholders if p.placeholder_format.type in TITLE_TYPES and not self._off_slide(p)), None)
        left, right = int(self.W * 0.06), int(self.W * 0.94)
        top = int(self.H * 0.08)
        if title is not None and title.has_text_frame and title.text_frame.text.strip():
            b = Box.of(title)
            left, right = max(left, b.left), min(right, b.left + b.width) if b.width > self.W * 0.4 else right
            top = b.top + b.height + self.gap
        bottom = int(self.H * 0.92)
        return Box(left, top, right - left, max(bottom - top, int(self.H * 0.3)))

    def _place_free(self, pslide, items: list[tuple[str, str]], content: dict, slot_types: dict, entry) -> None:
        """Content without a placeholder goes under the title: a kicker above it,
        short texts (captions) at the bottom, bodies to the left of visuals."""
        area = self._free_area(pslide)
        tops = [i for i in items if i[0] == "kicker" and i[1] == "text"]
        lines = [i for i in items if slot_kind(i[1], i[0]) in ("text", "title") and i not in tops]
        bodies = [i for i in items if slot_kind(i[1], i[0]) == "body" and i[1] != "cards"]
        visuals = [i for i in items if i not in tops + lines + bodies]
        line_h = int(self.H * 0.07)
        for n, (name, slot_type) in enumerate(tops):
            self._textbox(pslide, Box(area.left, area.top + n * line_h, area.width, line_h),
                          self._paras(slot_type, content[name]), size=Pt(16))
        area = Box(area.left, area.top + line_h * len(tops), area.width, area.height - line_h * (len(tops) + len(lines)))
        for n, (name, slot_type) in enumerate(lines):
            self._textbox(pslide, Box(area.left, area.top + area.height + n * line_h, area.width, line_h),
                          self._paras(slot_type, content[name]), size=Pt(16), anchor=MSO_ANCHOR.MIDDLE)
        if bodies and visuals:
            text_col, visual_col = area.columns([0.42, 0.58], self.gap * 2)
        else:
            text_col = visual_col = area
        entry.placed += [f"`{name}` under the title" for name, _ in items]
        for (name, slot_type), box in zip(bodies, text_col.rows(len(bodies), self.gap) if bodies else []):
            self._textbox(pslide, box, self._paras(slot_type, content[name]))
        for (name, slot_type), box in zip(visuals, visual_col.rows(len(visuals), self.gap) if visuals else []):
            slot = slot_types.get(name, {"type": slot_type})
            if slot_type == "image":
                self._image(pslide, content[name], box, entry, label=name)
            elif slot_type == "cards":
                self._cards_grid(pslide, content[name], slot.get("fields") or [], box, entry)
            else:
                self._visual(pslide, slot_type, content[name], box, entry, label=name)

    # -- overlays (D-013, D-016) ---------------------------------------------------------------------

    def _overlay(self, pslide, overlay: dict, entry) -> None:
        box = Box(int(overlay.get("x", 0) * self.W), int(overlay.get("y", 0) * self.H),
                  int(overlay.get("w", 0.2) * self.W), int(overlay.get("h", 0.1) * self.H))
        if "image" in overlay:
            self._image(pslide, {**overlay["image"], "fit": overlay["image"].get("fit", "contain")}, box, entry, label="extra image")
            return
        style = overlay.get("style") or {}
        pct = lambda v: int(self.W * float(v) / 100)  # noqa: E731 - lengths are % of the slide width
        has_box = style.get("background") or style.get("borderWidth")
        if style.get("radius") and has_box:
            shape = pslide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box.left, box.top, box.width, box.height)
            xml.set_corner_radius(shape, pct(style["radius"]) / max(1, min(box.width, box.height)))
        elif has_box:
            shape = pslide.shapes.add_shape(MSO_SHAPE.RECTANGLE, box.left, box.top, box.width, box.height)
        else:
            shape = pslide.shapes.add_textbox(box.left, box.top, box.width, box.height)
        if has_box:
            if style.get("background"):
                shape.fill.solid()
                shape.fill.fore_color.rgb = _rgb(style["background"])
                if style.get("backgroundOpacity", 1) < 1:
                    xml.set_fill_alpha(shape, style["backgroundOpacity"])
            else:
                shape.fill.background()
            if style.get("borderWidth"):
                shape.line.width = pct(style["borderWidth"])
                shape.line.color.rgb = _rgb(style.get("borderColor") or style.get("color") or "#000000")
                dash = {"dashed": MSO_LINE_DASH_STYLE.DASH, "dotted": MSO_LINE_DASH_STYLE.ROUND_DOT}.get(style.get("borderStyle"))
                if dash:
                    shape.line.dash_style = dash
            else:
                shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        pad = pct(style.get("padding", 0))
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = pad
        tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(style.get("valign"), MSO_ANCHOR.TOP)
        size = style.get("size", "body")
        size_pct = _NAMED_SIZES.get(size, 2.2) if isinstance(size, str) else float(size)
        paras = []
        for line in str(overlay.get("text", "")).split("\n"):
            paras.append(md.Para("p", 0, md.inline(line, self.glossary)))
        self._write_paras(shape, paras, size=Emu(int(self.W * size_pct / 100)), font=_FONTS.get(style.get("font", "body")))
        for p in tf.paragraphs:
            p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(style.get("align"), PP_ALIGN.LEFT)
            for r in p.runs:
                if style.get("bold"):
                    r.font.bold = True
                if style.get("italic"):
                    r.font.italic = True
                if style.get("underline"):
                    r.font.underline = True
                if style.get("uppercase"):
                    r.text = r.text.upper()
                if style.get("color"):
                    r.font.color.rgb = _rgb(style["color"])
                elif has_box and self.text_fill is None:  # a shape's own text would be white
                    r.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
                if style.get("outlineWidth"):
                    xml.set_text_outline(r, pct(style["outlineWidth"]) / 12700, style.get("outlineColor") or "#FFFFFF")

    # -- notes --------------------------------------------------------------------------------------

    def _notes(self, pslide, slide: dict) -> None:
        paras: list[md.Para] = []
        if self.include_notes and slide.get("notes"):
            paras += md.parse(slide["notes"], self.glossary)
        reader = slide.get("reader") or {}
        if self.include_reader and (reader.get("summary") or reader.get("blocks")):
            if paras:
                paras.append(md.Para("p", 0, []))
            paras.append(md.Para("p", 0, [md.Run("Reader layer", bold=True)]))
            if reader.get("summary"):
                paras.append(md.Para("p", 0, [md.Run(str(reader["summary"]), italic=True)]))
            for block in reader.get("blocks") or []:
                paras += self._block_paras(block)
        for extra in self.notes_extra:
            if paras:
                paras.append(md.Para("p", 0, []))
            paras += [md.Para("p", 0, [md.Run(line)]) for line in extra.split("\n")]
        if not paras:
            return
        notes = pslide.notes_slide.notes_placeholder
        if notes is None:
            return
        # Notes have no real bullets: list markers are written as text.
        flat = []
        for p in paras:
            if p.kind in ("p", "code"):
                flat.append(p)
            else:
                flat.append(md.Para("p", 0, [md.Run("    " * p.level + md.marker(p))] + p.runs))
        self._write_paras(notes, flat)

    def _block_paras(self, block: dict) -> list[md.Para]:
        kind = block.get("type")
        head = lambda text: [md.Para("p", 0, [md.Run(text, bold=True)])]  # noqa: E731
        body = lambda text: md.parse(text, self.glossary)  # noqa: E731
        if kind == "text":
            return body(block.get("md", ""))
        if kind == "details":
            return head(block.get("title", "")) + body(block.get("md", ""))
        if kind == "callout":
            return head(block.get("title") or block.get("variant", "Note").capitalize()) + body(block.get("md", ""))
        if kind == "tabs":
            out = []
            for tab in block.get("tabs") or []:
                out += head(tab.get("label", "")) + body(tab.get("md", ""))
            return out
        if kind == "image":
            img = block.get("image") or {}
            out = head("Picture: " + (img.get("alt") or img.get("src", "")))
            for h in block.get("hotspots") or []:
                out += [md.Para("bullet", 0, [md.Run(h.get("title", "") + ": " + md.plain(h.get("md", ""), self.glossary))])]
            return out
        if kind == "slide-ref":
            return [md.Para("p", 0, [md.Run("See also: " + (block.get("label") or block.get("slide", "")))])]
        if kind == "diagram":
            return head("Diagram (Mermaid)") + [md.Para("code", 0, [md.Run(line)]) for line in str(block.get("mermaid", "")).split("\n")]
        if kind == "table":
            labels, rows, _totals, _align, _ = datafiles.table_cells(block, self.read_file)
            return head(block.get("title") or "Table") + [md.Para("p", 0, [md.Run("\t".join(r))]) for r in [labels] + rows[:50]]
        if kind == "tag-cloud":
            terms, _ = datafiles.cloud_terms(block, self.read_file)
            return head(block.get("title") or "Tag cloud") + [md.Para("p", 0, [md.Run(", ".join(t for t, _ in terms))])]
        title = block.get("title") or _KIND_LABEL.get(kind, kind or "Block")
        where = block.get("url") or block.get("video") or block.get("pdf") or block.get("data") or ""
        return [md.Para("p", 0, [md.Run(f"{title}" + (f": {where}" if where else ""), bold=True)])]


def _has_value(value) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, (list, dict)):
        return bool(value)
    return True


def _rgb(color: str) -> RGBColor:
    c = color.lstrip("#")
    if len(c) == 3:
        c = "".join(ch * 2 for ch in c)
    return RGBColor.from_string(c.upper())


def _pick_level(styles: list, para: md.Para) -> int | None:
    """The outline level whose inherited style fits the paragraph: the n-th
    bullet level for a list item at depth n, the A. / 1. level for those
    lists, a plain level for text, the bold capitals level for a subheading.
    None when the placeholder has no such level (the bullet is then set on
    the paragraph itself)."""
    def plain(s) -> bool:
        return s.bullet in ("none", None) and not (s.bold or s.caps or s.serif)

    if para.kind in ("p", "code"):
        if para.role == "line" or (para.role == "heading" and styles[0].bullet in ("none", None)):
            return 0
        if para.is_subheading:
            sub = next((n for n, s in enumerate(styles) if s.bullet in ("none", None) and s.bold and s.caps), None)
            if sub is not None:
                return sub
        visible = [r for r in para.runs if r.text.strip()]
        if para.kind == "p" and visible and all(r.italic for r in visible):
            serif = next((n for n, s in enumerate(styles) if s.bullet in ("none", None) and s.serif), None)
            if serif is not None:
                return serif
        return next((n for n, s in enumerate(styles) if plain(s)), None)
    if para.kind == "bullet":
        chars = [n for n, s in enumerate(styles) if s.bullet == "char"]
        return chars[min(para.level, len(chars) - 1)] if chars else None
    return next((n for n, s in enumerate(styles) if s.bullet == para.kind), None)


def _svg_size(blob: bytes) -> tuple[float, float]:
    """An SVG's width and height (from width / height, else the viewBox); 4:3 when unknown."""
    head = blob[:4096].decode("utf-8", errors="replace")
    tag = re.search(r"<svg\b[^>]*>", head, re.S)
    attrs = tag.group(0) if tag else ""

    def num(name: str) -> float | None:
        m = re.search(rf'\b{name}\s*=\s*["\']\s*([\d.]+)\s*(px)?\s*["\']', attrs)
        return float(m.group(1)) if m else None

    w, h = num("width"), num("height")
    if w and h:
        return w, h
    box = re.search(r'viewBox\s*=\s*["\']\s*[-\d.]+[\s,]+[-\d.]+[\s,]+([\d.]+)[\s,]+([\d.]+)', attrs)
    if box and float(box.group(1)) and float(box.group(2)):
        return float(box.group(1)), float(box.group(2))
    return 400.0, 300.0


def _focus(focus) -> tuple[float, float]:
    """CSS object-position ('center', 'top left', '30% 70%') → fractions 0–1."""
    x = y = 0.5
    if not isinstance(focus, str):
        return x, y
    words = focus.lower().split()
    nums = [w for w in words if w.endswith("%")]
    if nums:
        x = float(nums[0][:-1]) / 100
        y = float(nums[1][:-1]) / 100 if len(nums) > 1 else 0.5
        return min(1, max(0, x)), min(1, max(0, y))
    if "left" in words:
        x = 0.0
    if "right" in words:
        x = 1.0
    if "top" in words:
        y = 0.0
    if "bottom" in words:
        y = 1.0
    return x, y
