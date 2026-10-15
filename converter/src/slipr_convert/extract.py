"""python-pptx objects → intermediate model (model.py).

What is extracted per slide: title, subtitle, text bodies (placeholders and
text boxes, in reading order), pictures with alt text, speaker notes, the
hidden flag, the section, build animations, tables, charts, SmartArt and
videos (I-121), the transition, entrance effects, zooms, the slide's own
background picture, and a list of elements that cannot be converted
(audio, OLE objects, newer chart kinds, …).

Pictures and shapes from the slide layout / master (logos, decorations) are
not part of `slide.shapes` and are therefore ignored, as intended: the Slipr
template provides the looks.
"""
from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from . import reader, text
from .model import ChartData, DeckModel, ImageRef, SlideModel, TextBody, Visual

_TITLE = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_SUBTITLE = {PP_PLACEHOLDER.SUBTITLE}
_IGNORED = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER}
_LIST_PLACEHOLDERS = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT}


def extract_deck(prs) -> DeckModel:
    section_of: dict[int, str] = {}
    section_names = []
    for name, ids in reader.sections(prs):
        section_names.append(name)
        for sid in ids:
            section_of[sid] = name

    slides = [extract_slide(slide, i, section_of.get(slide.slide_id)) for i, slide in enumerate(prs.slides)]
    # Section zooms open the first slide of their section.
    section_slides = reader.section_slides(prs)
    for sm, slide in zip(slides, prs.slides):
        for kind, target in reader.zoom_targets(slide):
            ids = [int(target)] if kind == "slide" and target.isdigit() else section_slides.get(target, [])[:1]
            sm.zoom_targets.extend(i for i in ids if i not in sm.zoom_targets)

    core = prs.core_properties
    stamp = core.modified or core.created
    ratio = (prs.slide_width or 4) / (prs.slide_height or 3)
    return DeckModel(
        title=(core.title or "").strip() or None,
        author=(core.author or "").strip() or None,
        date=stamp.date().isoformat() if stamp else None,
        aspect_ratio="4:3" if abs(ratio - 4 / 3) < 0.05 else "16:9",
        sections=section_names,
        slides=slides,
        slide_width=int(prs.slide_width or 0),
        slide_height=int(prs.slide_height or 0),
        slide_parts=[str(slide.part.partname) for slide in prs.slides],
    )


def extract_slide(slide, index: int, section: str | None) -> SlideModel:
    model = SlideModel(
        index=index,
        slide_id=slide.slide_id,
        layout_name=slide.slide_layout.name or "",
        hidden=reader.is_hidden(slide),
        section=section,
        has_build_animations=reader.has_build_animations(slide),
        transition=reader.transition(slide),
        entrances=reader.entrance_effects(slide),
    )
    if (background := reader.background_image(slide)) is not None:
        blob, ext, sha1 = background
        model.background = ImageRef(blob=blob, ext=ext, sha1=sha1, alt="")
    for shape in _walk(slide.shapes):
        _extract_shape(shape, model)

    model.bodies.sort(key=lambda b: (b.top, b.left))
    model.images.sort(key=lambda im: (im.top, im.left))
    model.visuals.sort(key=lambda v: (v.top, v.left))

    if slide.has_notes_slide:  # NB: accessing slide.notes_slide would create one
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            notes = text.notes_markdown(frame)
            model.notes = notes or None
    return model


def _walk(shapes):
    """All shapes, descending into groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def _extract_shape(shape, model: SlideModel) -> None:
    ph_type = shape.placeholder_format.type if shape.is_placeholder else None
    idx = shape.placeholder_format.idx if shape.is_placeholder else None
    if ph_type in _IGNORED:
        return
    top, left = int(shape.top or 0), int(shape.left or 0)

    # Videos (also in a placeholder, where the poster frame looks like a picture)
    media = reader.media_kind(shape)
    if media == "video" and (video := reader.video_of(shape)) is not None:
        model.visuals.append(Visual(kind="video", top=top, left=left, idx=idx, video=video))
        return
    if media is not None:
        model.unsupported.append("audio" if media == "audio" else "linked video")
        return

    # Pictures (plain pictures and filled picture placeholders)
    image = _image_of(shape)
    if image is not None:
        blob, ext, sha1 = image
        effects, opacity = reader.picture_effects(shape)
        model.images.append(ImageRef(
            blob=blob, ext=ext, sha1=sha1, alt=reader.alt_text(shape),
            top=top, left=left, width=int(shape.width or 0), height=int(shape.height or 0), idx=idx,
            effects=effects, opacity=opacity, crop=_crop_of(shape), shape_id=shape.shape_id, name=shape.name or "",
        ))
        return

    kind = reader.graphic_kind(shape)  # tables, charts, SmartArt, OLE (also inside placeholders)
    if kind == "table" and getattr(shape, "has_table", False):
        model.visuals.append(Visual(kind="table", top=top, left=left, idx=idx, rows=_table_rows(shape.table)))
        return
    if kind == "chart" and getattr(shape, "has_chart", False) and (chart := _chart_data(shape.chart)) is not None:
        model.visuals.append(Visual(kind="chart", top=top, left=left, idx=idx, chart=chart))
        return
    if kind == "smartArt" and (diagram := reader.smartart(shape)) is not None:
        if diagram.kind == "list":  # list SmartArt is just text: a Markdown list
            markdown = "\n".join(_list_lines(diagram.roots, 0))
            model.bodies.append(TextBody(markdown=markdown, is_list=True, placeholder="smartart", top=top, left=left, idx=idx,
                                         width=int(shape.width or 0), height=int(shape.height or 0)))
        else:
            model.visuals.append(Visual(kind="smartArt", top=top, left=left, idx=idx, diagram=diagram))
        return
    if shape.shape_type == MSO_SHAPE_TYPE.MEDIA or kind:
        model.unsupported.append(kind or "media")
        return

    if not getattr(shape, "has_text_frame", False) or not shape.text_frame.text.strip():
        return  # empty placeholder or a decorative shape

    if ph_type in _TITLE and model.title is None:
        model.title = text.plain_text(shape.text_frame)
        model.title_idx = idx
        return
    if ph_type in _SUBTITLE and model.subtitle is None:
        model.subtitle = text.plain_text(shape.text_frame)
        model.subtitle_idx = idx
        return

    as_list = ph_type in _LIST_PLACEHOLDERS or (shape.is_placeholder and ph_type is None)
    markdown = text.frame_markdown(shape.text_frame, as_list=as_list,
                                   style_of=lambda paragraph: reader.paragraph_style(shape, paragraph))
    if markdown:
        placeholder = ph_type.name.lower() if ph_type is not None else None
        # Text boxes and shapes keep their size and look: they may become text elements (I-089).
        look = None if shape.is_placeholder else reader.text_look(shape)
        model.bodies.append(TextBody(markdown=markdown, is_list=markdown.lstrip().startswith("- "),
                                     placeholder=placeholder, top=top, left=left, idx=idx,
                                     width=int(shape.width or 0), height=int(shape.height or 0), look=look,
                                     shape_id=shape.shape_id))


def _table_rows(table) -> list[list[str]]:
    """A table's cell texts, row by row (merged cells: the text stays in the first cell)."""
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cells.append("" if cell.is_spanned else " ".join(cell.text.split()))
        rows.append(cells)
    return rows


def _chart_data(chart) -> ChartData | None:
    """Categories × series of a chart, or None when it has no readable data."""
    try:
        plot = chart.plots[0]
        categories = [str(c) if c is not None else "" for c in plot.categories]
        series = [(s.name or f"Series {n}", [None if v is None else float(v) for v in s.values])
                  for n, s in enumerate(chart.series, 1)]
    except (IndexError, KeyError, TypeError, ValueError, AttributeError):
        return None
    if not series:
        return None
    if not categories:
        categories = [str(n) for n in range(1, max(len(v) for _, v in series) + 1)]
    title = None
    if chart.has_title and chart.chart_title.has_text_frame:
        title = " ".join(chart.chart_title.text_frame.text.split()) or None
    return ChartData(title=title, categories=categories, series=series, number_format=reader.chart_number_format(chart))


def _list_lines(nodes, depth: int) -> list[str]:
    lines = []
    for node in nodes:
        lines.append(f"{'  ' * min(depth, 3)}- {text.escape_markdown(node.text)}")
        lines.extend(_list_lines(node.children, depth + 1))
    return lines


def _crop_of(shape) -> tuple[float, float, float, float] | None:
    """The shown part of a cropped picture as (x, y, w, h), fractions of the
    picture, or None when it is not cropped. Negative crops (a picture
    "cropped outwards" in PowerPoint) count as 0."""
    try:
        left, top, right, bottom = (max(0.0, float(getattr(shape, f"crop_{side}") or 0.0))
                                    for side in ("left", "top", "right", "bottom"))
    except (AttributeError, TypeError, ValueError):
        return None
    w, h = 1.0 - left - right, 1.0 - top - bottom
    if w <= 0.01 or h <= 0.01 or (left < 0.0005 and top < 0.0005 and right < 0.0005 and bottom < 0.0005):
        return None
    return round(left, 4), round(top, 4), round(w, 4), round(h, 4)


def _image_of(shape):
    """(blob, ext, sha1) for shapes that carry a picture, else None.

    `shape.image` raises ValueError for linked (not embedded) pictures and
    for empty picture placeholders, so hasattr() cannot be used to probe it.
    """
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE and not shape.is_placeholder:
        return None
    try:
        img = shape.image
    except (AttributeError, ValueError, KeyError):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return None  # a linked picture: nothing to copy
        return None  # an empty picture placeholder, or not a picture at all
    return img.blob, (img.ext or "png").lower(), img.sha1
