"""Fixture decks, generated in code with python-pptx.

Uses python-pptx's built-in default template, whose layouts are named
"Title Slide", "Title and Content", "Section Header", "Two Content",
"Comparison", "Title Only", "Blank", "Content with Caption",
"Picture with Caption", the same names as PowerPoint's defaults.
"""
from __future__ import annotations

import io
import json
import struct
import uuid
import zlib
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[2]
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def png(rgb: tuple[int, int, int], w: int = 4, h: int = 3) -> bytes:
    """A tiny valid PNG of one colour."""
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def layout(prs, name: str):
    return next(l for l in prs.slide_layouts if l.name == name)


def placeholder(slide, *types):
    return next(ph for ph in slide.placeholders if ph.placeholder_format.type in types)


def set_alt(shape, text: str) -> None:
    for el in shape._element.iter(f"{{{P}}}cNvPr"):
        el.set("descr", text)
        return


A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def add_picture_effect(pic, effect: str) -> None:
    """Give a picture a PowerPoint effect, as PowerPoint writes it (I-084):
    'shadow', 'softEdges', 'roundedCorners' or 'border'."""
    sp_pr = pic._element.spPr
    if effect in ("shadow", "softEdges"):
        lst = sp_pr.find(f"{{{A}}}effectLst")
        if lst is None:
            lst = etree.SubElement(sp_pr, f"{{{A}}}effectLst")
        if effect == "shadow":
            shadow = etree.SubElement(lst, f"{{{A}}}outerShdw", blurRad="76200", dist="38100", dir="5400000")
            color = etree.SubElement(shadow, f"{{{A}}}srgbClr", val="000000")
            etree.SubElement(color, f"{{{A}}}alpha", val="40000")
        else:
            etree.SubElement(lst, f"{{{A}}}softEdge", rad="63500")
    elif effect == "roundedCorners":
        sp_pr.find(f"{{{A}}}prstGeom").set("prst", "roundRect")
    elif effect == "border":
        pic.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        pic.line.width = Inches(0.05)


def set_transparency(pic, percent: int) -> None:
    """PowerPoint's picture transparency: <a:alphaModFix amt> on the image."""
    blip = pic._element.find(f".//{{{A}}}blip")
    etree.SubElement(blip, f"{{{A}}}alphaModFix", amt=str((100 - percent) * 1000))


def add_build_animation(slide) -> None:
    """Minimal <p:timing> with a paragraph build, as PowerPoint writes for bullets."""
    body = next(sh for sh in slide.placeholders if sh.placeholder_format.type != PP_PLACEHOLDER.TITLE)
    spid = body.shape_id
    timing = etree.SubElement(slide._element, f"{{{P}}}timing")
    bld_lst = etree.SubElement(timing, f"{{{P}}}bldLst")
    etree.SubElement(bld_lst, f"{{{P}}}bldP", spid=str(spid), grpId="0", build="p")


def add_sections(prs, sections: list[tuple[str, list]]) -> None:
    root = prs.part._element
    ext_lst = root.find(f"{{{P}}}extLst")
    if ext_lst is None:
        ext_lst = etree.SubElement(root, f"{{{P}}}extLst")
    ext = etree.SubElement(ext_lst, f"{{{P}}}ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    sec_lst = etree.SubElement(ext, f"{{{P14}}}sectionLst", nsmap={"p14": P14})
    for name, slides in sections:
        sec = etree.SubElement(sec_lst, f"{{{P14}}}section", name=name, id="{" + str(uuid.uuid4()).upper() + "}")
        ids = etree.SubElement(sec, f"{{{P14}}}sldIdLst")
        for slide in slides:
            etree.SubElement(ids, f"{{{P14}}}sldId", id=str(slide.slide_id))


def build_sample_deck() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    prs.core_properties.title = "Quarterly Review"
    prs.core_properties.author = "Test Author"

    # 1. Title slide
    s1 = prs.slides.add_slide(layout(prs, "Title Slide"))
    s1.shapes.title.text = "Quarterly Review"
    placeholder(s1, PP_PLACEHOLDER.SUBTITLE).text = "Q3 2026"

    # 2. Bullets with levels, bold, italic, a link, notes, build animation
    s2 = prs.slides.add_slide(layout(prs, "Title and Content"))
    s2.shapes.title.text = "Highlights"
    tf = placeholder(s2, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY).text_frame
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = "Revenue "
    r = p.add_run(); r.text = "up"; r.font.bold = True
    p2 = tf.add_paragraph(); p2.level = 1
    r = p2.add_run(); r.text = "in "
    r = p2.add_run(); r.text = "all"; r.font.italic = True
    r = p2.add_run(); r.text = " regions"
    p3 = tf.add_paragraph()
    r = p3.add_run(); r.text = "See report"; r.hyperlink.address = "https://example.com/report"
    p4 = tf.add_paragraph(); p4.text = "Costs *stable* [flat]"  # literal Markdown characters
    s2.notes_slide.notes_text_frame.text = "Mention the new team."
    add_build_animation(s2)
    # …a "New" badge: a filled rounded rectangle with white capitals (→ a text element, I-089)
    badge = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10), Inches(0.5), Inches(2), Inches(0.6))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(0xBC, 0x44, 0x8C)
    badge.line.fill.background()
    badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = badge.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "New"
    r.font.size = Pt(14); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r._r.get_or_add_rPr().set("cap", "all")  # noqa: SLF001
    # …and a shape in the theme's default style (accent 1 fill, white text, dark outline)
    draft = s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10), Inches(6), Inches(2), Inches(1))
    draft.text_frame.text = "Draft"

    # 3. Section header (its body has no bullets via the layout)
    s3 = prs.slides.add_slide(layout(prs, "Section Header"))
    s3.shapes.title.text = "Details"
    placeholder(s3, PP_PLACEHOLDER.BODY).text = "The numbers behind it"

    # 4. Two content
    s4 = prs.slides.add_slide(layout(prs, "Two Content"))
    s4.shapes.title.text = "Before and after"
    bodies = [ph for ph in s4.placeholders if ph.placeholder_format.type != PP_PLACEHOLDER.TITLE]
    bodies[0].text_frame.text = "Manual work"
    bodies[1].text_frame.text = "Automated"

    # 5. Picture with caption (picture placeholder, alt text)
    s5 = prs.slides.add_slide(layout(prs, "Picture with Caption"))
    s5.shapes.title.text = "Our office"
    pic = placeholder(s5, PP_PLACEHOLDER.PICTURE).insert_picture(io.BytesIO(png((200, 30, 30))))
    set_alt(pic, "The office building at sunrise")
    add_picture_effect(pic, "shadow")      # → picture style "shadow" (I-084)
    set_transparency(pic, 30)              # → opacity 0.7
    placeholder(s5, PP_PLACEHOLDER.BODY).text = "Amsterdam, 2026"

    # 6. Blank slide with a text box and a picture without alt text
    s6 = prs.slides.add_slide(layout(prs, "Blank"))
    box = s6.shapes.add_textbox(Inches(1), Inches(6), Inches(6), Inches(1))
    box.text_frame.text = "Free text box"
    s6.shapes.add_picture(io.BytesIO(png((30, 30, 200))), Inches(1), Inches(1), Inches(4), Inches(3))
    # …and a second picture that fits no image slot: an extra image at its place (I-083)
    logo = s6.shapes.add_picture(io.BytesIO(png((30, 160, 60))), Inches(7.5), Inches(5.5), Inches(2), Inches(1.5))
    set_alt(logo, "Company logo")
    logo.crop_left, logo.crop_right, logo.crop_top = 0.1, 0.2, 0.25  # → crop x 0.1, w 0.7, y 0.25, h 0.75 (I-090)

    # 7. A table: amounts, Dutch percentages and a totals row (→ a table slot, I-121)
    s7 = prs.slides.add_slide(layout(prs, "Title Only"))
    s7.shapes.title.text = "A table"
    table = s7.shapes.add_table(4, 3, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, cells in enumerate([["Region", "Revenue", "Growth"], ["North", "€ 1.200", "12,5%"],
                               ["South", "€ 800", "-5%"], ["Total", "€ 2.000", ""]]):
        for c, value in enumerate(cells):
            table.cell(r, c).text = value

    # 8. A hidden backup slide
    s8 = prs.slides.add_slide(layout(prs, "Title and Content"))
    s8.shapes.title.text = "Backup"
    placeholder(s8, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY).text = "Extra numbers"
    s8._element.set("show", "0")

    add_sections(prs, [("Intro", [s1, s2]), ("Main part", [s3, s4, s5, s6, s7, s8])])
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# -- Fixtures for I-121: XML that python-pptx cannot write itself -------------------------------------

R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
PSLZ = "http://schemas.microsoft.com/office/powerpoint/2016/slidezoom"


def set_transition(slide, name: str) -> None:
    """A slide transition as PowerPoint writes it; Morph inside mc:AlternateContent."""
    from pptx.oxml import parse_xml

    if name == "morph":
        xml = (f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="{P}"><mc:Choice xmlns:p159="{P159}" Requires="p159">'
               f'<p:transition spd="slow"><p159:morph option="byObject"/></p:transition></mc:Choice>'
               f'<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    else:
        xml = f'<p:transition xmlns:p="{P}" spd="med"><p:{name}/></p:transition>'
    slide._element.find(f"{{{P}}}clrMapOvr").addnext(parse_xml(xml))


def add_entrance(slide, shape, preset: int, subtype: int = 0) -> None:
    """A click-triggered entrance effect on a whole shape (p:timing, simplified)."""
    from pptx.oxml import parse_xml

    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="3" fill="hold"><p:childTnLst><p:par><p:cTn id="4" presetID="{preset}" presetClass="entr" '
           f'presetSubtype="{subtype}" fill="hold" nodeType="clickEffect"><p:childTnLst><p:set><p:cBhvr>'
           f'<p:cTn id="5" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{shape.shape_id}"/></p:tgtEl>'
           f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
           f'<p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    slide._element.append(parse_xml(xml))


def add_smartart(slide, layout_id: str, nodes: list, n: int) -> None:
    """A SmartArt graphic: its data part (points and parOf connections) and the graphic frame.
    `nodes`: [(text, [children…]), …]."""
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml import parse_xml

    model = etree.Element(f"{{{DGM}}}dataModel", nsmap={"dgm": DGM, "a": A})
    pts = etree.SubElement(model, f"{{{DGM}}}ptLst")
    cxns = etree.SubElement(model, f"{{{DGM}}}cxnLst")
    doc = etree.SubElement(pts, f"{{{DGM}}}pt", modelId="0", type="doc")
    etree.SubElement(doc, f"{{{DGM}}}prSet", loTypeId=f"urn:microsoft.com/office/officeart/2005/8/layout/{layout_id}")
    counter = iter(range(1, 1000))

    def add(parent: str, items: list) -> None:
        for order, (text, children) in enumerate(items):
            mid = str(next(counter))
            pt = etree.SubElement(pts, f"{{{DGM}}}pt", modelId=mid)
            t = etree.SubElement(pt, f"{{{DGM}}}t")
            etree.SubElement(t, f"{{{A}}}bodyPr")
            run = etree.SubElement(etree.SubElement(t, f"{{{A}}}p"), f"{{{A}}}r")
            etree.SubElement(run, f"{{{A}}}t").text = text
            etree.SubElement(pts, f"{{{DGM}}}pt", modelId=f"pres{mid}", type="pres")  # presentation points are skipped
            etree.SubElement(cxns, f"{{{DGM}}}cxn", modelId=f"c{mid}", srcId=parent, destId=mid, srcOrd=str(order), destOrd="0")
            etree.SubElement(cxns, f"{{{DGM}}}cxn", modelId=f"p{mid}", type="presOf", srcId=mid, destId=f"pres{mid}", srcOrd="0", destOrd="0")
            add(mid, children)

    add("0", nodes)
    part = Part(PackURI(f"/ppt/diagrams/data{n}.xml"), "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
                slide.part.package, etree.tostring(model, xml_declaration=True, encoding="UTF-8", standalone=True))
    rid = slide.part.relate_to(part, RT.DIAGRAM_DATA)
    frame = parse_xml(
        f'<p:graphicFrame xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"><p:nvGraphicFramePr><p:cNvPr id="{100 + n}" name="Diagram {n}"/>'
        f'<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="914400" y="1828800"/><a:ext cx="7315200" cy="3657600"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
        f'<dgm:relIds xmlns:dgm="{DGM}" r:dm="{rid}" r:lo="" r:qs="" r:cs=""/></a:graphicData></a:graphic></p:graphicFrame>')
    slide.shapes._spTree.append(frame)  # noqa: SLF001


def add_slide_zoom(slide, target) -> None:
    """A slide zoom to `target`, as PowerPoint writes it (mc:AlternateContent; python-pptx doesn't list it)."""
    from pptx.oxml import parse_xml

    xml = (f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="{P}" xmlns:a="{A}"><mc:Choice xmlns:pslz="{PSLZ}" Requires="pslz">'
           f'<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="200" name="Slide Zoom"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
           f'<p:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="514350"/></p:xfrm><a:graphic><a:graphicData uri="{PSLZ}">'
           f'<pslz:sldZm><pslz:sldZmObj sldId="{target.slide_id}" cId="0"/></pslz:sldZm></a:graphicData></a:graphic></p:graphicFrame>'
           f'</mc:Choice><mc:Fallback/></mc:AlternateContent>')
    slide.shapes._spTree.append(parse_xml(xml))  # noqa: SLF001


def add_slide_link(slide, run, target) -> None:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    rid = slide.part.relate_to(target.part, RT.SLIDE)
    link = etree.SubElement(run._r.get_or_add_rPr(), f"{{{A}}}hlinkClick", action="ppaction://hlinksldjump")  # noqa: SLF001
    link.set(f"{{{R}}}id", rid)


def set_background_picture(slide, blob: bytes) -> None:
    from pptx.oxml import parse_xml

    _part, rid = slide.part.get_or_add_image_part(io.BytesIO(blob))
    bg = parse_xml(f'<p:bg xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"><p:bgPr><a:blipFill><a:blip r:embed="{rid}"/>'
                   f'<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    slide._element.cSld.insert(0, bg)  # noqa: SLF001


def build_feature_deck() -> bytes:
    """One slide per I-121 feature: chart, video, online video, SmartArt
    (process, hierarchy, list), a slide link and a zoom, entrance effects,
    a background picture, transitions and Morph."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "Features"

    # 1. A chart → a table with its data (fade)
    s1 = prs.slides.add_slide(layout(prs, "Title Only"))
    s1.shapes.title.text = "Sales chart"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3"]
    data.add_series("2025", (10, 12.5, 14))
    data.add_series("2026", (11, 13, 18))
    chart = s1.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(8), Inches(4), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Sales per quarter"
    set_transition(s1, "fade")

    # 2. An embedded video with a poster, trim and bookmarks (push)
    s2 = prs.slides.add_slide(layout(prs, "Title Only"))
    s2.shapes.title.text = "Launch video"
    movie = s2.shapes.add_movie(io.BytesIO(b"\x00\x00\x00\x18ftypmp42 fake video"), Inches(1), Inches(2), Inches(8), Inches(4.5),
                                poster_frame_image=io.BytesIO(png((10, 10, 10))), mime_type="video/mp4")
    set_alt(movie, "The launch")
    media = movie._element.find(f".//{{{P14}}}media")  # noqa: SLF001
    etree.SubElement(media, f"{{{P14}}}trim", st="1500", end="2000")
    bmks = etree.SubElement(media, f"{{{P14}}}bmkLst")
    etree.SubElement(bmks, f"{{{P14}}}bmk", name="Lift-off", time="4200")
    etree.SubElement(bmks, f"{{{P14}}}bmk", name="Countdown", time="0")
    set_transition(s2, "push")

    # 3. An online video (YouTube): an external link instead of an embedded file (push)
    s3 = prs.slides.add_slide(layout(prs, "Title Only"))
    s3.shapes.title.text = "Online video"
    online = s3.shapes.add_movie(io.BytesIO(b"placeholder"), Inches(1), Inches(2), Inches(8), Inches(4.5), mime_type="video/mp4")
    set_alt(online, "Keynote")
    ext = online._element.find(f".//{{{P14}}}media").getparent()  # noqa: SLF001
    ext.getparent().remove(ext)
    rid = s3.part.relate_to("https://www.youtube.com/embed/abcDEF12345?rel=0", RT.VIDEO, is_external=True)
    online._element.find(f".//{{{A}}}videoFile").set(f"{{{R}}}link", rid)  # noqa: SLF001
    set_transition(s3, "push")

    # 4. A process SmartArt → a left-to-right Mermaid flowchart
    s4 = prs.slides.add_slide(layout(prs, "Title Only"))
    s4.shapes.title.text = "Process"
    add_smartart(s4, "process1", [("Plan", []), ("Build", [("Test", [])]), ("Ship", [])], 1)

    # 5. A hierarchy SmartArt → a top-down flowchart
    s5 = prs.slides.add_slide(layout(prs, "Title Only"))
    s5.shapes.title.text = "Organisation"
    add_smartart(s5, "orgChart1", [("CEO", [("CFO", []), ("CTO \"tech\"", [])])], 2)

    # 6. A link to slide 1 and a "next slide" jump in the text; a zoom to slide 2
    s6 = prs.slides.add_slide(layout(prs, "Title and Content"))
    s6.shapes.title.text = "Links"
    p = placeholder(s6, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY).text_frame.paragraphs[0]
    r = p.add_run(); r.text = "See "
    r = p.add_run(); r.text = "the chart"
    add_slide_link(s6, r, s1)
    r = p.add_run(); r.text = " or "
    r = p.add_run(); r.text = "go on"
    etree.SubElement(r._r.get_or_add_rPr(), f"{{{A}}}hlinkClick", action="ppaction://hlinkshowjump?jump=nextslide")  # noqa: SLF001
    add_slide_zoom(s6, s2)

    # 7. A photo, a "!!logo" that flies in from below, and a background picture
    photo = png((120, 140, 160), 8, 6)
    s7 = prs.slides.add_slide(layout(prs, "Blank"))
    pic = s7.shapes.add_picture(io.BytesIO(photo), Inches(1), Inches(1), Inches(6), Inches(4.5))
    set_alt(pic, "A photo")
    logo = s7.shapes.add_picture(io.BytesIO(png((200, 0, 0))), Inches(9), Inches(5.5), Inches(2), Inches(1))
    set_alt(logo, "Logo")
    logo.name = "!!logo"
    add_entrance(s7, logo, 2, 4)
    set_background_picture(s7, png((240, 240, 230)))

    # 8. The same photo and logo elsewhere, with a Morph transition
    s8 = prs.slides.add_slide(layout(prs, "Blank"))
    pic = s8.shapes.add_picture(io.BytesIO(photo), Inches(1), Inches(1), Inches(6), Inches(4.5))
    set_alt(pic, "A photo")
    logo = s8.shapes.add_picture(io.BytesIO(png((200, 0, 0))), Inches(1), Inches(6), Inches(3), Inches(1.2))
    set_alt(logo, "Logo")
    logo.name = "!!logo"
    set_transition(s8, "morph")

    # 9. A list SmartArt → a Markdown list
    s9 = prs.slides.add_slide(layout(prs, "Title Only"))
    s9.shapes.title.text = "Values"
    add_smartart(s9, "vList2", [("Open", [("Honest", [])]), ("Curious", [])], 3)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


@pytest.fixture(scope="session")
def sample_deck() -> bytes:
    return build_sample_deck()


@pytest.fixture(scope="session")
def feature_deck() -> bytes:
    return build_feature_deck()


def js_validate(show: dict, template_name: str, tmp_path: Path) -> dict:
    """Run shared/validate.js on a show: {errors, warnings} (the contract with the viewer)."""
    import subprocess

    show_file = tmp_path / "show.json"
    show_file.write_text(json.dumps(show), encoding="utf-8")
    script = (
        "import { validateShow } from " + json.dumps((REPO / "shared" / "validate.js").as_uri()) + ";"
        "import { readFileSync } from 'node:fs';"
        f"const show = JSON.parse(readFileSync({json.dumps(str(show_file))}, 'utf8'));"
        f"const template = JSON.parse(readFileSync({json.dumps(str(REPO / 'templates' / template_name / 'template.json'))}, 'utf8'));"
        "const r = validateShow(show, { template });"
        "console.log(JSON.stringify({ errors: r.errors, warnings: r.warnings }));"
    )
    proc = subprocess.run(["node", "--input-type=module", "-e", script], capture_output=True, text=True, check=True)
    return json.loads(proc.stdout)


@pytest.fixture(scope="session")
def template() -> dict:
    return json.loads((REPO / "templates" / "default" / "template.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="session")
def converted(sample_deck, template):
    from slipr_convert import convert

    return convert(sample_deck, template, template_ref="../../templates/default", show_id="sample", source_name="sample.pptx")
