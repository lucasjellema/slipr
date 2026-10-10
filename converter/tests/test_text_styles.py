"""Paragraph styles → Slipr Markdown (letter / numbered lists, subheadings, serif lines)."""
from __future__ import annotations

import io

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from slipr_convert import convert

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def styled(paragraph, *, autonum=None, no_bullet=False, bold=False, caps=False, font=None):
    """Give a paragraph explicit properties, as a template's outline levels would."""
    ppr = paragraph._p.get_or_add_pPr()
    if no_bullet:
        etree.SubElement(ppr, f"{{{A}}}buNone")
    if autonum:
        etree.SubElement(ppr, f"{{{A}}}buAutoNum", type=autonum)
    rpr = etree.SubElement(ppr, f"{{{A}}}defRPr")
    if bold:
        rpr.set("b", "1")
    if caps:
        rpr.set("cap", "all")
    if font:
        etree.SubElement(rpr, f"{{{A}}}latin", typeface=font)


def test_template_text_styles_become_markdown(template):
    prs = Presentation()
    slide = prs.slides.add_slide(next(l for l in prs.slide_layouts if l.name == "Title and Content"))
    slide.shapes.title.text = "Styles"
    body = next(ph for ph in slide.placeholders if ph.placeholder_format.type != PP_PLACEHOLDER.TITLE)
    tf = body.text_frame
    tf.paragraphs[0].text = "Bullet"
    p = tf.add_paragraph(); p.text = "Intro line"; styled(p, no_bullet=True)
    p = tf.add_paragraph(); p.text = "Subheading"; styled(p, no_bullet=True, bold=True, caps=True)
    p = tf.add_paragraph(); p.text = "Serif line"; styled(p, no_bullet=True, font="Source Serif 4 48pt")
    p = tf.add_paragraph(); p.text = "Letter one"; styled(p, autonum="alphaUcPeriod")
    p = tf.add_paragraph(); p.text = "Letter two"; styled(p, autonum="alphaUcPeriod")
    p = tf.add_paragraph(); p.text = "Number one"; styled(p, autonum="arabicPeriod")
    p = tf.add_paragraph(); p.text = "Deep"; p.level = 6  # deeper levels nest at most 4 deep
    out = io.BytesIO()
    prs.save(out)

    show = convert(out.getvalue(), template, template_ref="t").show
    assert show["slides"][0]["content"]["body"] == (
        "- Bullet\n\n"
        "Intro line\n\n"
        "**Subheading**\n\n"
        "*Serif line*\n\n"
        "A. Letter one\nA. Letter two\n\n"
        "1. Number one\n\n"
        "      - Deep"
    )
