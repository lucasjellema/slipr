"""The reverse direction (I-045): a Slipr show → PowerPoint, checked by reading the deck back."""
from __future__ import annotations

import io
import json

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from slipr_convert import convert
from slipr_convert.export import export_show
from slipr_convert.export import markdown as md
from slipr_convert.export_cli import main
from .conftest import P14, REPO

ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"


def load(name: str) -> tuple[dict, dict, callable]:
    folder = REPO / "shows" / name
    show = json.loads((folder / "show.json").read_text(encoding="utf-8"))
    template_name = show["template"]["ref"].rstrip("/").rsplit("/", 1)[-1]
    template = json.loads((REPO / "templates" / template_name / "template.json").read_text(encoding="utf-8"))

    def read_file(rel: str) -> bytes | None:
        path = folder / rel
        return path.read_bytes() if path.is_file() else None

    return show, template, read_file


@pytest.fixture(scope="module")
def demo():
    show, template, read_file = load("demo")
    result = export_show(show, template, read_file=read_file)
    return show, result, Presentation(io.BytesIO(result.pptx))


def texts(slide) -> str:
    return "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def notes(slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


# -- Markdown ---------------------------------------------------------------------------------------

def test_markdown_blocks():
    paras = md.parse("Intro line\nsecond line\n\n- one\n  - nested\n- two\n\n3. three\n4. four\n\nB. bee\n\n```\ncode()\n```")
    assert [(p.kind, p.level) for p in paras] == [("p", 0), ("bullet", 0), ("bullet", 1), ("bullet", 0),
                                                  ("number", 0), ("number", 0), ("upper", 0), ("code", 0)]
    assert paras[0].runs[1].is_break
    assert paras[4].start == 3 and paras[6].start == 2
    assert md.parse(["a", "b"])[1].text == "b"


def test_markdown_inline():
    runs = md.inline("**bold** *it* `x` [site](https://example.com) [[reveal-js]] [[t|own]]^[popup] \\*lit\\* [slide](#s-02)",
                     {"reveal-js": {"term": "Reveal.js"}})
    assert [(r.text, r.bold, r.italic, r.code, r.link) for r in runs[:6]] == [
        ("bold", True, False, False, None), (" ", False, False, False, None), ("it", False, True, False, None),
        (" ", False, False, False, None), ("x", False, False, True, None), (" ", False, False, False, None)]
    assert runs[6].link == "https://example.com"
    assert "".join(r.text for r in runs[7:]) == " Reveal.js own *lit* slide"
    assert md.parse("**Key findings**")[0].is_subheading
    assert md.plain("- a\n- b") == "• a\n• b"


# -- the demo show, on python-pptx's default layouts --------------------------------------------------

def test_demo_exports_every_slide(demo):
    show, result, prs = demo
    assert len(prs.slides) == len(show["slides"]) == len(result.report.slides)
    assert prs.slide_width == 12192000 and prs.slide_height == 6858000  # 16:9
    assert result.report.skipped_count == 0
    assert prs.core_properties.title == show["meta"]["title"]
    assert prs.slides[1].shapes.title.text == "Why another presentation tool?"


def test_layouts_and_placeholders(demo):
    show, _result, prs = demo
    layouts = [s.slide_layout.name for s in prs.slides]
    assert layouts[0] == "Title Slide"
    assert layouts[2] == "Comparison"  # two-column with column headings: the mapping that fits best
    for slide in prs.slides:  # no "Click to add text" prompts left
        for ph in slide.placeholders:
            assert ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE or not ph.has_text_frame or ph.text_frame.text.strip()
    body = prs.slides[1].placeholders[1].text_frame
    assert body.paragraphs[0].runs[-1].font.bold  # "**room**"
    assert [p.level for p in body.paragraphs] == [0, 0, 1, 0]


def test_sections_hidden_slides_and_notes(demo):
    show, _result, prs = demo
    root = prs.part._element  # noqa: SLF001
    names = [s.get("name") for s in root.iter(f"{{{P14}}}section")]
    assert names == [c["title"] for c in show["chapters"]]
    ids = [int(e.get("id")) for e in root.iter(f"{{{P14}}}sldId")]
    assert ids == [s.slide_id for s in prs.slides]  # every slide in a section, in order
    hidden = [i for i, s in enumerate(show["slides"]) if (s.get("audience") or {}).get("hidden")]
    assert hidden and all(prs.slides[i]._element.get("show") == "0" for i in hidden)  # noqa: SLF001
    problem = notes(prs.slides[1])
    assert "Reader layer" in problem and show["slides"][1]["reader"]["summary"] in problem


def test_rich_slots(demo):
    show, result, prs = demo
    by_id = {s["id"]: prs.slides[i] for i, s in enumerate(show["slides"])}
    table = next(sh for sh in by_id["s-table"].shapes if sh.has_table).table
    assert table.cell(0, 0).text == "Chapter" and "Total" in table.cell(len(table.rows) - 1, 0).text
    assert "templates" in texts(by_id["s-tag-cloud"])
    assert "Amsterdam" in texts(by_id["s-map"])
    assert "Mermaid" in notes(by_id["s-pipeline"]) or any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in by_id["s-pipeline"].shapes)
    assert any(sh.shape_type == MSO_SHAPE_TYPE.MEDIA for sh in by_id["s-video"].shapes)  # a real video
    assert any("Graph" in f or "graph" in f for f in (f for s in result.report.slides for f in s.fallbacks))


def test_svg_pictures_keep_their_svg(demo):
    show, _result, prs = demo
    slide = prs.slides[[s["id"] for s in show["slides"]].index("s-architecture")]
    blips = [b for sh in slide.shapes for b in sh._element.iter(f"{{{ASVG}}}svgBlip")]  # noqa: SLF001
    assert blips
    rid = blips[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    assert slide.part.related_part(rid).content_type == "image/svg+xml"


def test_overlays(demo):
    show, _result, prs = demo
    i = [s["id"] for s in show["slides"]].index("s-tailored")
    box = next(sh for sh in prs.slides[i].shapes if "many" in (sh.text_frame.text if sh.has_text_frame else ""))
    overlay = show["slides"][i]["overlays"][0]
    assert abs(box.left / prs.slide_width - overlay["x"]) < 0.01
    assert box.fill.fore_color.rgb is not None
    assert any(r.font.italic for p in box.text_frame.paragraphs for r in p.runs)  # "*many*"


def test_notes_and_reader_layer_can_be_left_out():
    show, template, read_file = load("demo")
    prs = Presentation(io.BytesIO(export_show(show, template, read_file=read_file, include_notes=False, include_reader=False).pptx))
    assert all(not notes(s).strip() or "Mermaid" in notes(s) for s in prs.slides)


def test_conclusion_demo_without_base():
    show, template, read_file = load("conclusion-demo")
    result = export_show(show, template, read_file=read_file)
    prs = Presentation(io.BytesIO(result.pptx))
    assert len(prs.slides) == len(show["slides"])
    assert result.report.skipped_count == 0
    cards = show["slides"][[s["layout"] for s in show["slides"]].index("columns-4")]["content"]["cards"]
    slide = prs.slides[[s["layout"] for s in show["slides"]].index("columns-4")]
    assert all(c["title"] in texts(slide) for c in cards)


# -- a base deck with template roles --------------------------------------------------------------------

def base_deck() -> bytes:
    prs = Presentation()
    by_name = {layout.name: layout for layout in prs.slide_layouts}
    sample = prs.slides.add_slide(by_name["Title and Content"])
    sample.shapes.title.text = "An old slide"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_base_deck_and_roles():
    template = {"layouts": [{"id": "statement", "slots": [{"name": "title", "type": "text"}, {"name": "kicker", "type": "text"},
                                                          {"name": "body", "type": "richtext"}]}],
                "pptxMapping": {"Comparison": {"layout": "statement", "placeholders": {"0": "title", "1": "kicker", "4": "body"}}}}
    show = {"slipr": "0.1", "id": "t", "meta": {"title": "T"}, "template": {"ref": "x"},
            "slides": [{"id": "s-1", "layout": "statement", "content": {"title": "Hello", "kicker": "Intro", "body": "Some **text**"}}]}
    prs = Presentation(io.BytesIO(export_show(show, template, read_file=lambda _: None, base=base_deck()).pptx))
    assert len(prs.slides) == 1  # the base deck's own slide is gone
    slide = prs.slides[0]
    assert slide.slide_layout.name == "Comparison"
    by_idx = {p.placeholder_format.idx: p for p in slide.placeholders}
    assert by_idx[0].text_frame.text == "Hello"
    assert by_idx[1].text_frame.text == "Intro"
    assert by_idx[4].text_frame.text == "Some text"
    assert 2 not in by_idx and 3 not in by_idx  # empty placeholders removed


# -- round trip ------------------------------------------------------------------------------------------

def test_round_trip_through_the_converter(sample_deck, template):
    converted = convert(sample_deck, template, template_ref="../../templates/default")

    def read_file(rel: str) -> bytes | None:
        return next((a.data for a in converted.assets if a.path == rel), None)

    exported = export_show(converted.show, template, read_file=read_file)
    again = convert(exported.pptx, template, template_ref="../../templates/default")
    titles = [s["content"].get("title") for s in converted.show["slides"]]
    assert [s["content"].get("title") for s in again.show["slides"]] == titles
    assert [s["layout"] for s in again.show["slides"]] == [s["layout"] for s in converted.show["slides"]]


# -- command line ----------------------------------------------------------------------------------------

def test_cli(tmp_path, capsys):
    out = tmp_path / "demo.pptx"
    assert main([str(REPO / "shows" / "demo"), "--out", str(out)]) == 0
    assert out.is_file() and (tmp_path / "demo-export-report.md").is_file()
    assert "Exported 23 slides" in capsys.readouterr().out
    assert main([str(REPO / "shows" / "demo"), "--out", str(out)]) == 1
    assert "--force" in capsys.readouterr().err
    assert main([str(REPO / "shows" / "demo" / "show.json"), "--out", str(out), "--force", "--no-reader"]) == 0
    assert main([str(tmp_path / "nothing"), "--out", str(out)]) == 1


def test_svg_part_is_well_formed(demo):
    _show, result, _prs = demo
    import zipfile
    with zipfile.ZipFile(io.BytesIO(result.pptx)) as z:
        types = etree.fromstring(z.read("[Content_Types].xml"))
        svg_parts = [n for n in z.namelist() if n.endswith(".svg")]
        assert svg_parts
        declared = {e.get("PartName", "").lstrip("/") for e in types} | {e.get("Extension") for e in types}
        assert all(p in declared or "svg" in declared for p in svg_parts)
