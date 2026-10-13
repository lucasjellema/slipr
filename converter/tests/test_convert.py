"""End-to-end: a generated deck through convert(), checked slide by slide."""
from __future__ import annotations

from slipr_convert.validate import validate_show


def slide(converted, number: int) -> dict:
    return converted.show["slides"][number - 1]


def test_output_validates_against_the_shared_schema(converted):
    assert validate_show(converted.show) == []
    assert converted.report.warnings == []


def test_meta_settings_and_chapters(converted):
    show = converted.show
    assert show["id"] == "sample"
    assert show["meta"]["title"] == "Quarterly Review"
    assert show["meta"]["author"] == "Test Author"
    assert len(show["meta"]["date"]) == 10
    assert show["settings"]["aspectRatio"] == "16:9"
    assert show["template"] == {"ref": "../../templates/default", "version": "^1.4"}
    assert show["chapters"] == [{"id": "intro", "title": "Intro"}, {"id": "main-part", "title": "Main part"}]
    assert [s["id"] for s in show["slides"]] == [f"s-0{n}" for n in range(1, 9)]
    assert slide(converted, 1)["chapter"] == "intro"
    assert slide(converted, 3)["chapter"] == "main-part"


def test_title_slide(converted):
    s = slide(converted, 1)
    assert s["layout"] == "title"
    assert s["content"] == {"title": "Quarterly Review", "subtitle": "Q3 2026"}


def test_bullets_formatting_notes_and_builds(converted):
    s = slide(converted, 2)
    assert s["layout"] == "bullets"
    assert s["content"]["title"] == "Highlights"
    assert s["content"]["body"] == (
        "- Revenue **up**\n"
        "  - in *all* regions\n"
        "- [See report](https://example.com/report)\n"
        "- Costs \\*stable\\* \\[flat\\]"
    )
    assert s["notes"] == "Mention the new team."
    assert s["audience"] == {"fragments": True}


def test_section_header_uses_inherited_no_bullets(converted):
    s = slide(converted, 3)
    assert s["layout"] == "section"
    assert s["content"] == {"title": "Details", "subtitle": "The numbers behind it"}
    assert "reader" not in s


def test_two_content(converted):
    s = slide(converted, 4)
    assert s["layout"] == "two-column"
    assert s["content"]["left"] == "- Manual work"
    assert s["content"]["right"] == "- Automated"


def test_picture_with_caption(converted):
    s = slide(converted, 5)
    assert s["layout"] == "image-full"
    # The picture's shadow becomes the template's "shadow" style, its 30 % transparency opacity 0.7 (I-084).
    assert s["content"]["image"] == {"src": "images/slide05-1.png", "alt": "The office building at sunrise", "style": "shadow", "opacity": 0.7}
    assert s["content"]["caption"] == "Amsterdam, 2026"


def test_blank_slide_text_box_and_picture_without_alt(converted):
    s = slide(converted, 6)
    assert s["layout"] == "image-full"
    assert s["content"]["image"]["src"] == "images/slide06-1.png"
    assert s["content"]["caption"] == "Free text box"
    entry = converted.report.slides[5]
    assert any("no alt text" in n for n in entry.notes)
    assert any("no title" in n for n in entry.notes)


def test_picture_without_a_slot_becomes_an_extra_image(converted):
    s = slide(converted, 6)
    # 7.5 x 5.5 in, 2 x 1.5 in on python-pptx's 13.333 x 7.5 in (16:9) slide
    assert s["overlays"] == [{"image": {"src": "images/slide06-2.png", "alt": "Company logo", "crop": {"x": 0.1, "y": 0.25, "w": 0.7, "h": 0.75}},
                              "x": 0.5625, "y": 0.7333, "w": 0.15, "h": 0.2}]
    assert list(s).index("overlays") == list(s).index("content") + 1
    assert "reader" not in s, "nothing moved to the reader layer"
    assert converted.report.slides[5].placed == ["picture 2"]
    assert converted.show["slipr"] == "0.2", "extra images need format 0.2"
    assert "Placed as an extra image: picture 2" in converted.report.to_markdown()


def test_unsupported_table_is_reported_and_flagged_in_notes(converted):
    s = slide(converted, 7)
    assert s["layout"] == "bullets"  # template pptxMapping: "Title Only" → bullets
    assert "table" in s["notes"]
    assert converted.report.slides[6].skipped == ["table"]


def test_hidden_slide_becomes_backup(converted):
    s = slide(converted, 8)
    assert s["audience"] == {"hidden": True}
    assert any("backup" in n for n in converted.report.slides[7].notes)


def test_assets(converted):
    assert [a.path for a in converted.assets] == ["images/slide05-1.png", "images/slide06-1.png", "images/slide06-2.png"]
    assert all(a.data.startswith(b"\x89PNG") for a in converted.assets)


def test_report_markdown(converted):
    md = converted.report.to_markdown()
    assert md.startswith("# Conversion report")
    assert "### 2. Highlights → `s-02`" in md
    assert "⚠ Not converted: table" in md


def test_notes_options(sample_deck, template):
    from slipr_convert import convert

    without = convert(sample_deck, template, template_ref="t", include_notes=False)
    assert "notes" not in without.show["slides"][1]
    to_reader = convert(sample_deck, template, template_ref="t", notes_to_reader=True)
    assert to_reader.show["slides"][1]["reader"]["blocks"] == [{"id": "notes", "type": "text", "md": "Mention the new team."}]


def test_picture_effects_map_to_template_styles(template):
    """Soft edges, rounded corners and a border map to the template's styles;
    with several effects the first in reader.PICTURE_EFFECTS wins (I-084)."""
    import io

    from pptx import Presentation
    from pptx.util import Inches

    from slipr_convert import reader
    from slipr_convert.convert import convert
    from .conftest import add_picture_effect, png

    prs = Presentation()
    blank = next(l for l in prs.slide_layouts if l.name == "Blank")
    pics = []
    for effects in (["softEdges"], ["roundedCorners"], ["border"], ["border", "shadow"], []):
        s = prs.slides.add_slide(blank)
        pic = s.shapes.add_picture(io.BytesIO(png((90, 90, 90))), Inches(1), Inches(1), Inches(4), Inches(3))
        for e in effects:
            add_picture_effect(pic, e)
        pics.append(pic)
    assert [reader.picture_effects(p)[0] for p in pics] == [["softEdges"], ["roundedCorners"], ["border"], ["shadow", "border"], []]

    out = io.BytesIO()
    prs.save(out)
    result = convert(out.getvalue(), template, template_ref="../../templates/default", show_id="fx", source_name="fx.pptx")
    styles = [s["content"]["image"].get("style") for s in result.show["slides"]]
    assert styles == ["soft", "rounded", "frame", "shadow", None]
    assert result.show["slipr"] == "0.2"

    # A template without picture styles: no style, and the report says why.
    plain = {k: v for k, v in template.items() if k != "pictureStyles"}
    result = convert(out.getvalue(), plain, template_ref="../../templates/default", show_id="fx", source_name="fx.pptx")
    assert all("style" not in s["content"]["image"] for s in result.show["slides"])
    assert result.show["slipr"] == "0.1"
    assert any("no picture style for softEdges" in n for n in result.report.slides[0].notes)
